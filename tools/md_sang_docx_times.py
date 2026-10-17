#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Chuyển tài liệu Markdown sang Word (.docx) — FONT MẶC ĐỊNH TIMES NEW ROMAN.

VÌ SAO CÓ (22/08/2026):
Repo đã có `EBM-Dashboards/tools/build_dashboard_docx.py`, nhưng công cụ đó chỉ
đọc được khối `DATA` của một dashboard Evidence Workbench — nó KHÔNG chuyển được
một tài liệu Markdown thường (báo cáo, hồ sơ kiểm kê, bản giao kiến trúc). Trước
nay các tài liệu đó chỉ có bản `.md`, mà `.md` thì không mở được bằng Word và
không in ra giấy cho hội đồng/đồng nghiệp đọc được.

NGUYÊN TẮC ĐI CÙNG DÂY CHUYỀN SẴN CÓ: bản `.docx` sinh TỪ CHÍNH file `.md` nguồn,
không gõ lại tay và không dựng lại nội dung từ dữ liệu khác — cùng lý do mà bộ năm
xuất đồng thời của dashboard tồn tại: hai bản dựng từ hai nguồn sẽ lệch nhau mà
không ai nhận ra.

FONT: Times New Roman toàn văn (thân bài · đề mục · bảng · trích dẫn), khớp quy ước
`build_dashboard_docx.py` đã dùng cho mọi tài liệu lưu trữ của hệ. Đoạn mã/đường dẫn
giữ NGUYÊN Times New Roman nhưng tô nền xám nhạt để vẫn phân biệt được — không đổi
sang font khác, vì mặc định đã được ấn định là Times New Roman.

PHẠM VI CỐ Ý HẸP — công cụ này KHÔNG kiểm nội dung y khoa, KHÔNG chấm chất lượng,
KHÔNG thay `verify_dashboard.py` hay bất kỳ cổng nào. Nó chỉ đổi định dạng.

Cú pháp Markdown hỗ trợ: đề mục `#`–`####` · đoạn văn · **đậm** · *nghiêng* ·
`mã` · [chữ](liên-kết) · bảng dạng ống · trích dẫn `>` · danh sách `-`/`*`/`1.` ·
đường kẻ ngang `---`.

CHẾ ĐỘ TRÌNH BÀY ĐỀ CƯƠNG (`--trang-de-cuong`, thêm 15/10/2026): dùng cho văn bản nộp
Hội đồng (đề cương, thuyết minh). Khổ A4; lề trái 3 cm, phải/trên/dưới 2 cm; đoạn văn
căn đều hai bên, giãn dòng 1,3; số trang giữa chân trang. Đường kẻ `---` KHÔNG vẽ
thành đường kẻ mà thành NGẮT TRANG; mọi nội dung đứng trước `---` đầu tiên là TRANG BÌA
và được căn giữa.

CHUẨN TRÌNH BÀY TÀI LIỆU NỘP (vá 17/10/2026): bản đầu của chế độ đề cương để bảng ở cỡ
thân − 1 (13 ⇒ 12pt) và giữ nguyên ký tự trang trí («☐» U+2610 — Times New Roman không có
hình chữ, in ra ô vuông rỗng), nên trục ④ của `kiem_chi_tiet_he_nghien_cuu.py` báo đề cương
nộp Hội đồng của C1a sai chuẩn (đo 17/10: bảng 12pt, 267 ký tự trang trí). Nguyên nhân gốc:
công cụ này nằm ở repo gốc, là một điểm `doc.save` KHÔNG đi qua `chuan_trinh_bay.py` như
11 bộ sinh .docx của cổng G0–G10. Nay chế độ đề cương gọi thẳng
`medical-ebm-automation/tools/chuan_trinh_bay.py::ap_dinh_dang_tai_lieu` (nguồn DUY NHẤT
của bảng ký tự — không chép bảng sang đây): Times New Roman, thân `--co-chu`, bảng 11pt,
«☐» → «[ ]», «☑» → «[x]»…; lề, căn đều, ngắt trang, trang bìa, cấp bậc đề mục và số trang
giữ nguyên. Không nạp được repo y khoa ⇒ bảng vẫn 11pt nhưng ký tự trang trí CHƯA được
làm sạch: công cụ vẫn ghi tệp, báo rõ và trả mã 1 (không in dấu xanh giả).
Chế độ thường không đổi.

Dùng:
    python3 tools/md_sang_docx_times.py <đầu-vào>.md [--ra <đầu-ra>.docx]
    python3 tools/md_sang_docx_times.py bao-cao.md --co-chu 12
    python3 tools/md_sang_docx_times.py de-cuong.md --co-chu 13 --trang-de-cuong

Mã thoát: 0 = xong · 1 = đã ghi .docx nhưng CHƯA làm sạch ký tự trang trí (chỉ ở
`--trang-de-cuong`, khi không nạp được `chuan_trinh_bay.py`) · 2 = lỗi đầu vào hoặc thiếu
python-docx.
"""
from __future__ import annotations

import argparse
import functools
import importlib.util
import re
import sys
from pathlib import Path

for _luong in (sys.stdout, sys.stderr):  # Windows cp1252 không in được «☐»/dấu tiếng Việt
    if hasattr(_luong, "reconfigure"):
        _luong.reconfigure(encoding="utf-8")

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover - phụ thuộc môi trường
    print("❌ Thiếu python-docx. Cài: pip install python-docx "
          "(venv ~/.ebm-venv đã có sẵn).", file=sys.stderr)
    raise SystemExit(2) from None

FONT = "Times New Roman"
XAM_MA = "EEEEEE"       # nền đoạn mã
XAM_CHU = RGBColor(0x44, 0x44, 0x44)
# Cỡ chữ trong bảng ở chế độ đề cương — PHẢI bằng chuan_trinh_bay.CO_CHU_BANG (test canh lệch).
# Chỉ để lối dự phòng khi không nạp được repo y khoa vẫn ra bảng đúng cỡ; có repo thì
# ap_dinh_dang_tai_lieu tự đặt theo hằng của chính nó.
CO_CHU_BANG_DE_CUONG = 11
GOC_REPO = Path(__file__).resolve().parents[1]

# Inline: mã trước tiên để nội dung bên trong không bị parse tiếp.
_INLINE = re.compile(
    r"(`[^`]+`)"
    r"|(\*\*[^*]+\*\*)"
    r"|(\*[^*\n]+\*)"
    r"|(\[[^\]]+\]\([^)]+\))"
)


def _dat_font(run, dam=False, nghieng=False, ma=False, co: float | None = None):
    run.font.name = FONT
    # Word cần khai riêng cho bảng mã Đông Á, nếu không dấu tiếng Việt có thể
    # rơi về font thay thế trên một số máy.
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for thuoc_tinh in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(thuoc_tinh), FONT)
    run.bold = dam
    run.italic = nghieng
    if co:
        run.font.size = Pt(co)
    if ma:
        run.font.color.rgb = XAM_CHU
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), XAM_MA)
        rpr.append(shd)


def _viet_inline(para, text: str, co: float | None = None, dam_het=False):
    """Ghi một đoạn văn có định dạng nội tuyến vào paragraph."""
    vi_tri = 0
    for m in _INLINE.finditer(text):
        if m.start() > vi_tri:
            _dat_font(para.add_run(text[vi_tri:m.start()]), dam=dam_het, co=co)
        tho = m.group(0)
        if tho.startswith("`"):
            _dat_font(para.add_run(tho[1:-1]), ma=True, co=co)
        elif tho.startswith("**"):
            _dat_font(para.add_run(tho[2:-2]), dam=True, co=co)
        elif tho.startswith("["):
            chu = tho[1:tho.index("]")]
            lien_ket = tho[tho.index("(") + 1:-1]
            _dat_font(para.add_run(chu), dam=dam_het, co=co)
            # Giữ URL dạng chữ để bản in giấy vẫn tra được nguồn.
            if lien_ket.startswith("http"):
                _dat_font(para.add_run(f" ({lien_ket})"), ma=True,
                          co=(co - 1.5) if co else 9)
        else:
            _dat_font(para.add_run(tho[1:-1]), nghieng=True, co=co)
        vi_tri = m.end()
    if vi_tri < len(text):
        _dat_font(para.add_run(text[vi_tri:]), dam=dam_het, co=co)


def _o_bang(row) -> list[str]:
    tho = row.strip().strip("|")
    return [c.strip() for c in re.split(r"(?<!\\)\|", tho)]


def _la_dong_ngan_cach(dong: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", dong)) and "-" in dong


def _so_trang_chan_trang(section) -> None:
    """Chèn trường PAGE căn giữa ở chân trang."""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for loai in ("begin", "instr", "end"):
        if loai == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), loai)
        run._r.append(el)
    _dat_font(run)


def _trinh_bay_de_cuong(doc) -> None:
    """Khổ A4, lề 3-2-2-2 cm, số trang; đoạn văn thân bài căn đều, giãn dòng 1,3."""
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin, s.right_margin = Cm(3), Cm(2)
        _so_trang_chan_trang(s)
    for p in doc.paragraphs:
        if p.alignment is not None or p.style.name.startswith("Heading"):
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.3
    for bang in doc.tables:
        for hang in bang.rows:
            for o in hang.cells:
                for p in o.paragraphs:
                    p.paragraph_format.line_spacing = 1.1
                    p.paragraph_format.space_after = Pt(2)


def _nap_tep(ten: str, duong: Path):
    """Nạp một module theo ĐƯỜNG DẪN TỆP — không chèn sys.path (tránh che module trùng tên)."""
    spec = importlib.util.spec_from_file_location(ten, duong)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@functools.lru_cache(maxsize=1)
def _nap_chuan_trinh_bay():
    """→ (module chuan_trinh_bay | None, đường dẫn tệp hoặc lý do không nạp được).

    `chuan_trinh_bay.py` của repo y khoa là nguồn DUY NHẤT của chuẩn trình bày tài liệu
    nộp (font · cỡ thân/bảng · bảng ký tự trang trí) mà trục ④ của
    `kiem_chi_tiet_he_nghien_cuu.py` đo theo. Repo y khoa dò qua `ban_sao_tran.duong_goc()`
    (vị trí lồng ở máy thật, anh em ở phiên Cloud) — không tự ghép đường dẫn.
    """
    try:
        bst = _nap_tep("_md_docx_ban_sao_tran", Path(__file__).resolve().parent / "ban_sao_tran.py")
        goc = bst.duong_goc("medical-ebm-automation", GOC_REPO)
    except (OSError, ImportError, AttributeError) as e:
        return None, f"không dò được repo y khoa ({type(e).__name__}: {e})"
    if goc is None:
        return None, "không thấy medical-ebm-automation/ (lồng hoặc anh em với repo gốc)"
    tep = goc / "tools" / "chuan_trinh_bay.py"
    if not tep.is_file():
        return None, f"thiếu {tep.as_posix()}"
    try:
        return _nap_tep("_md_docx_chuan_trinh_bay", tep), tep.as_posix()
    except Exception as e:  # noqa: BLE001 — mọi lỗi nạp đều phải thành lời báo, không sập công cụ
        return None, f"nạp {tep.as_posix()} lỗi ({type(e).__name__}: {e})"


def _ap_chuan_trinh_bay(doc, co_chu: float) -> bool:
    """Áp chuẩn trình bày tài liệu nộp lên văn bản đề cương; True nếu đã áp đủ.

    Gọi `chuan_trinh_bay.ap_dinh_dang_tai_lieu` — cùng điểm mà 11 bộ sinh .docx của cổng
    G0–G10 đi qua và đúng chuẩn mà `kiem_docx` đo: Times New Roman, thân `co_chu`, bảng
    `CO_CHU_BANG`, ký tự trang trí thay bằng văn bản («☐» → «[ ]», «☑» → «[x]»…). Hàm đó chỉ
    đụng font/cỡ/chữ của run trong thân bài và bảng — lề, căn đều, ngắt trang, căn giữa
    trang bìa và trường số trang ở chân trang giữ nguyên.

    Mũi tên đổi theo ĐÚNG bộ làm sạch chung như mọi cổng («→» → « đến », «⇒» → « suy ra »…)
    — bác sĩ chọn 17/10/2026, không giữ riêng cho đề cương.

    Một chỗ bù có chủ ý: hàm đó ghi cỡ thân bài lên MỌI run, kể cả đề mục, nên tên đề cương
    (thân + 6) và đề mục (thân + 1…3) sẽ phẳng thành cỡ thân. Run nào bộ dựng đã đặt LỚN HƠN
    cỡ thân (chỉ có ở đề mục và trang bìa) được trả lại đúng cỡ cũ để giữ cấp bậc đề mục;
    trục ④ đếm cỡ theo số ký tự nên vài dòng đề mục không đổi kết luận về thân bài.

    Không nạp được chuan_trinh_bay ⇒ False: bảng vẫn 11pt (đặt lúc dựng) nhưng ký tự trang
    trí còn nguyên — nơi gọi phải báo, không được coi là đạt.
    """
    ctb, _ = _nap_chuan_trinh_bay()
    if ctb is None:
        return False
    nguong = Pt(co_chu)
    giu_co = [(r, r.font.size) for p in doc.paragraphs for r in p.runs
              if r.font.size is not None and r.font.size > nguong]
    ctb.ap_dinh_dang_tai_lieu(doc, co_chu=co_chu)
    for r, co in giu_co:
        r.font.size = co
    return True


def dung_docx(md: str, co_chu: float, trang_de_cuong: bool = False) -> "Document":
    doc = Document()
    # Trạng thái chế độ đề cương: đang ở trang bìa? đoạn kế tiếp có ngắt trang trước?
    trang = {"bia": trang_de_cuong, "ngat": False}

    def _moi(p):
        """Áp trạng thái trang bìa/ngắt trang cho đoạn vừa tạo."""
        if trang["ngat"]:
            p.paragraph_format.page_break_before = True
            trang["ngat"] = False
        if trang["bia"]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
        return p

    # Font mặc định cho TOÀN tài liệu.
    for ten_style in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
                      "Heading 4", "List Bullet", "List Number", "Quote"):
        try:
            st = doc.styles[ten_style]
        except KeyError:
            continue
        st.font.name = FONT
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for tt in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(tt), FONT)
        if ten_style == "Normal":
            st.font.size = Pt(co_chu)
            st.font.color.rgb = RGBColor(0, 0, 0)

    dong_md = md.splitlines()
    i = 0
    while i < len(dong_md):
        dong = dong_md[i]
        tho = dong.strip()

        if not tho:
            i += 1
            continue

        # Đường kẻ ngang
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", tho):
            if trang_de_cuong:
                # Chế độ đề cương: đường kẻ = ngắt trang; hết trang bìa ở đường kẻ đầu tiên.
                trang["bia"] = False
                trang["ngat"] = True
                i += 1
                continue
            p = doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bot = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single")
            bot.set(qn("w:sz"), "6")
            bot.set(qn("w:color"), "BBBBBB")
            bdr.append(bot)
            ppr.append(bdr)
            i += 1
            continue

        # Đề mục
        m = re.match(r"^(#{1,4})\s+(.*)$", tho)
        if m:
            muc = len(m.group(1))
            chu = m.group(2).strip()
            if muc == 1:
                p = _moi(doc.add_paragraph())
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _viet_inline(p, chu, co=co_chu + 6, dam_het=True)
            else:
                p = _moi(doc.add_heading(level=min(muc, 4)))
                p.text = ""
                _viet_inline(p, chu, co=co_chu + (4 - muc))
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0x14, 0x25, 0x2C)
                    if trang["bia"]:
                        r.font.size = Pt(co_chu + 3)
                        r.bold = True
            i += 1
            continue

        # Bảng
        if tho.startswith("|") and i + 1 < len(dong_md) and _la_dong_ngan_cach(dong_md[i + 1]):
            dau = _o_bang(dong)
            i += 2
            than = []
            while i < len(dong_md) and dong_md[i].strip().startswith("|"):
                than.append(_o_bang(dong_md[i]))
                i += 1
            if trang["ngat"]:
                _moi(doc.add_paragraph())
            co_bang = CO_CHU_BANG_DE_CUONG if trang_de_cuong else co_chu - 1
            bang = doc.add_table(rows=1, cols=len(dau))
            bang.style = "Table Grid"
            for j, o in enumerate(dau):
                cell = bang.rows[0].cells[j]
                cell.text = ""
                _viet_inline(cell.paragraphs[0], o, co=co_bang, dam_het=True)
            for hang in than:
                cells = bang.add_row().cells
                for j, o in enumerate(hang[:len(dau)]):
                    cells[j].text = ""
                    _viet_inline(cells[j].paragraphs[0], o, co=co_bang)
            doc.add_paragraph()
            continue

        # Trích dẫn
        if tho.startswith(">"):
            khoi = []
            while i < len(dong_md) and dong_md[i].strip().startswith(">"):
                khoi.append(dong_md[i].strip().lstrip(">").strip())
                i += 1
            p = _moi(doc.add_paragraph())
            p.paragraph_format.left_indent = Pt(24)
            _viet_inline(p, " ".join(khoi), co=co_chu)
            for r in p.runs:
                r.italic = True
            continue

        # Danh sách
        m = re.match(r"^(\s*)([-*•]|\d+[.)])\s+(.*)$", dong)
        if m:
            thut = len(m.group(1)) // 2
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            noi_dung = [m.group(3).strip()]
            i += 1
            # gộp dòng tiếp nối (thụt lề, không phải mục mới)
            while (i < len(dong_md) and dong_md[i].strip()
                   and not re.match(r"^(\s*)([-*•]|\d+[.)])\s+", dong_md[i])
                   and dong_md[i].startswith(("  ", "\t"))
                   and not dong_md[i].strip().startswith(("#", "|", ">"))):
                noi_dung.append(dong_md[i].strip())
                i += 1
            try:
                p = _moi(doc.add_paragraph(style=style))
            except KeyError:
                p = _moi(doc.add_paragraph())
            p.paragraph_format.left_indent = Pt(18 + 18 * thut)
            _viet_inline(p, " ".join(noi_dung), co=co_chu)
            continue

        # Đoạn văn thường — gộp các dòng liền nhau
        khoi = [tho]
        i += 1
        while (i < len(dong_md) and dong_md[i].strip()
               and not dong_md[i].strip().startswith(("#", "|", ">", "---"))
               and not re.match(r"^(\s*)([-*•]|\d+[.)])\s+", dong_md[i])):
            khoi.append(dong_md[i].strip())
            i += 1
        p = _moi(doc.add_paragraph())
        if not trang["bia"]:
            p.paragraph_format.space_after = Pt(6)
        _viet_inline(p, " ".join(khoi), co=co_chu)

    if trang_de_cuong:
        _trinh_bay_de_cuong(doc)
        _ap_chuan_trinh_bay(doc, co_chu)
    return doc


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description="Markdown → Word (.docx), font mặc định Times New Roman.")
    ap.add_argument("nguon", type=Path, help="file .md đầu vào")
    ap.add_argument("--ra", type=Path, help="file .docx đầu ra (mặc định cùng tên)")
    ap.add_argument("--co-chu", type=float, default=12.0, help="cỡ chữ thân bài (pt)")
    ap.add_argument("--trang-de-cuong", action="store_true",
                    help="trình bày văn bản nộp Hội đồng: A4, lề 3-2-2-2 cm, căn đều, "
                         "«---» = ngắt trang, phần trước «---» đầu tiên là trang bìa căn giữa")
    a = ap.parse_args(argv)

    if not a.nguon.exists():
        print(f"❌ Không tìm thấy {a.nguon}", file=sys.stderr)
        return 2
    ra = a.ra or a.nguon.with_suffix(".docx")
    ra.parent.mkdir(parents=True, exist_ok=True)

    doc = dung_docx(a.nguon.read_text(encoding="utf-8"), a.co_chu, a.trang_de_cuong)
    doc.save(str(ra))
    if a.trang_de_cuong:
        ctb, nguon = _nap_chuan_trinh_bay()
        if ctb is None:
            print(f"🟡 {ra}  ({ra.stat().st_size:,} byte · font {FONT} · {a.co_chu:g}pt · "
                  f"bảng {CO_CHU_BANG_DE_CUONG}pt) — CHƯA làm sạch ký tự trang trí («☐»…): {nguon}. "
                  "Bản này CHƯA đạt chuẩn nộp; chạy lại trên máy có medical-ebm-automation/.",
                  file=sys.stderr)
            return 1
        print(f"✅ {ra}  ({ra.stat().st_size:,} byte · font {FONT} · {a.co_chu:g}pt · "
              f"bảng {ctb.CO_CHU_BANG}pt · chuẩn trình bày: {nguon})")
        return 0
    print(f"✅ {ra}  ({ra.stat().st_size:,} byte · font {FONT} · {a.co_chu:g}pt)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
