# -*- coding: utf-8 -*-
"""Chế độ đề cương của `md_sang_docx_times.py` phải ra đúng chuẩn trình bày tài liệu nộp (17/10/2026).

Bối cảnh: đề cương nộp Hội đồng của C1a dựng bằng `--trang-de-cuong` ra bảng 12pt (cỡ thân − 1) và giữ 267 ký tự
«☐» (Times New Roman không có hình chữ) — trục ④ của `medical-ebm-automation/tools/kiem_chi_tiet_he_nghien_cuu.py`
báo sai chuẩn. Test đo bằng CHÍNH `kiem_docx` của trục ④ (chạy ở tiến trình con để module repo y khoa không chen vào
sys.path của bộ test gốc), và đối chiếu với bản dựng KHÔNG áp chuẩn để chứng minh: thước đo không mù (bản thô có ký
tự lạ), lề · khổ giấy · căn lề · ngắt trang · số trang không đổi, và chữ đổi ĐÚNG như bộ làm sạch chung
`lam_sach_van_ban` — kể cả mũi tên «→» thành « đến » như mọi cổng (bác sĩ chọn 17/10/2026). Ngoại tuyến hoàn toàn, văn
bản fixture tự soạn.

Không có repo y khoa: bản sao git trần (CI/cloud) ⇒ skip có lý do; máy thật mà không nạp được ⇒ ĐỎ (cùng quy ước
`tools/conftest.py::_bo_qua_nhom_repo_y_khoa`).
"""
from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

pytest.importorskip("docx")
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

REPO = Path(__file__).resolve().parents[1]


def _nap(ten: str, duong: Path):
    spec = importlib.util.spec_from_file_location(ten, duong)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


M = _nap("_md_docx_ctb_test", REPO / "tools/md_sang_docx_times.py")

MD = """**ĐƠN VỊ THỬ**

# ĐỀ CƯƠNG THỬ

---

## 1. PHIẾU KHẢO SÁT

Câu 1. Anh/chị có thẻ bảo hiểm y tế không? ☐ Có ☐ Không. Câu 2. Anh/chị đã ký phiếu đồng thuận chưa? ☐ Rồi ☐ Chưa.

Sơ đồ nhân quả giả định: Tuổi → BHYT → thời gian chờ → hài lòng; đồng thời học vấn (→ thời gian chờ và → hài lòng).

- ☑ Người tham gia đã được giải thích mục tiêu, quyền rút lui và cách bảo mật thông tin trước khi trả lời phiếu.

---

| Lựa chọn | Đánh dấu |
|---|---|
| Đồng ý | ☐ |
| → Từ chối | ☑ |

### 1.1. Mục con

Đoạn văn thân bài dài hơn để cỡ thân bài chiếm ưu thế khi đếm theo số ký tự trong tài liệu thử nghiệm này.
"""


def _ctb_hoac_bo_qua():
    ctb, ly_do = M._nap_chuan_trinh_bay()
    if ctb is not None:
        return ctb
    bst = _nap("_bst_md_docx_ctb_test", REPO / "tools/ban_sao_tran.py")
    if bst.ban_sao_git_tran(REPO):
        pytest.skip(f"bản sao git trần — không có repo y khoa để nạp chuan_trinh_bay/kiem_docx: {ly_do}")
    pytest.fail(f"máy thật (còn gốc dữ liệu OneDrive) mà không nạp được chuan_trinh_bay: {ly_do}")


def _kiem_docx(ctb, tep: Path) -> dict:
    """Đo bằng kiem_chi_tiet_he_nghien_cuu.kiem_docx của repo y khoa, trong tiến trình con riêng."""
    goc = Path(ctb.__file__).resolve().parents[1]
    ma = ("import json, sys; from pathlib import Path; sys.path.insert(0, sys.argv[1]); "
          "import kiem_chi_tiet_he_nghien_cuu as K; print(json.dumps(K.kiem_docx(Path(sys.argv[2]))))")
    kq = subprocess.run([sys.executable, "-B", "-c", ma, str(goc / "tools"), str(tep)], cwd=str(goc),
                        capture_output=True, text=True, encoding="utf-8", timeout=180)
    assert kq.returncode == 0, kq.stderr[-2000:]
    return json.loads(kq.stdout.strip().splitlines()[-1])


def _van_ban(doc) -> list[str]:
    out = [p.text for p in doc.paragraphs]
    out += [p.text for t in doc.tables for h in t.rows for o in h.cells for p in o.paragraphs]
    return out


def _bo_cuc(doc) -> dict:
    s = doc.sections[0]
    return {
        "kho_le": [int(x) for x in (s.page_width, s.page_height, s.left_margin, s.right_margin,
                                    s.top_margin, s.bottom_margin)],
        "ngat_trang": [i for i, p in enumerate(doc.paragraphs) if p.paragraph_format.page_break_before],
        "can_le": [None if p.alignment is None else int(p.alignment) for p in doc.paragraphs],
        "so_trang": [e.get(qn("w:fldCharType")) for e in s.footer.paragraphs[0]._p.iter(qn("w:fldChar"))],
    }


def _ban_tho(monkeypatch, tmp_path: Path) -> Path:
    """Dựng bản KHÔNG áp chuẩn (giả lập vắng repo y khoa) — làm mốc so sánh."""
    with monkeypatch.context() as mp:
        mp.setattr(M, "_nap_chuan_trinh_bay", lambda: (None, "giả lập vắng repo y khoa"))
        md = tmp_path / "tho" / "DE_CUONG_NOP_HOI_DONG_THU.md"
        md.parent.mkdir()
        md.write_text(MD, encoding="utf-8")
        assert M.main([str(md), "--co-chu", "13", "--trang-de-cuong"]) == 1
    return md.with_suffix(".docx")


def test_kiem_docx_do_dat_chuan_va_giu_bo_cuc(tmp_path, monkeypatch):
    ctb = _ctb_hoac_bo_qua()
    tho = _ban_tho(monkeypatch, tmp_path)
    md = tmp_path / "DE_CUONG_NOP_HOI_DONG_THU.md"
    md.write_text(MD, encoding="utf-8")
    assert M.main([str(md), "--co-chu", "13", "--trang-de-cuong"]) == 0
    dat = md.with_suffix(".docx")

    d = _kiem_docx(ctb, dat)
    assert d["font_uu_the"] == "Times New Roman" and d["font_la"] == {}, d
    assert d["co_than_uu_the"] == 13.0, d
    assert d["co_bang_uu_the"] == float(ctb.CO_CHU_BANG) == 11.0, d
    assert d["ky_tu_la"] == 0, d

    # Thước đo không mù: chính bản thô (không áp chuẩn) bị trục ④ đếm ra ký tự trang trí.
    assert _kiem_docx(ctb, tho)["ky_tu_la"] > 0

    doc_dat, doc_tho = Document(dat), Document(tho)
    assert _bo_cuc(doc_dat) == _bo_cuc(doc_tho), "lề/khổ/căn lề/ngắt trang/số trang phải giữ nguyên"
    assert len(_bo_cuc(doc_dat)["ngat_trang"]) == 2 and _bo_cuc(doc_dat)["so_trang"] == ["begin", "end"]

    # Chữ đổi đúng như bộ làm sạch chung (mỗi đoạn fixture là một run), không thêm không bớt.
    assert _van_ban(doc_dat) == [ctb.lam_sach_van_ban(t) for t in _van_ban(doc_tho)]
    chu = "\n".join(_van_ban(doc_dat))
    assert "☐" not in chu and "☑" not in chu
    assert "[ ] Có [ ] Không" in chu and "[x] Người tham gia" in chu
    # Mũi tên theo chuẩn chung như mọi cổng — không giữ riêng cho đề cương.
    assert "→" not in chu and "Tuổi đến BHYT đến thời gian chờ đến hài lòng" in chu

    # Cấp bậc đề mục giữ nguyên: tên đề cương vẫn thân + 6.
    ten = next(p for p in doc_dat.paragraphs if p.text == "ĐỀ CƯƠNG THỬ")
    assert {r.font.size for r in ten.runs} == {Pt(19)}


def test_hang_co_bang_khop_chuan_trinh_bay():
    assert M.CO_CHU_BANG_DE_CUONG == _ctb_hoac_bo_qua().CO_CHU_BANG


def test_du_phong_khong_repo_y_khoa_bang_11_va_bao_ma_1(tmp_path, monkeypatch, capsys):
    """Chạy được ở mọi nơi có python-docx: vắng repo y khoa ⇒ bảng vẫn 11pt, ký tự trang trí CÒN (không tự chép
    bảng ký tự), công cụ vẫn ghi tệp nhưng trả mã 1 và báo rõ — không in dấu xanh."""
    tep = _ban_tho(monkeypatch, tmp_path)
    loi = capsys.readouterr()
    assert tep.exists()
    assert "CHƯA làm sạch" in loi.err and "✅" not in loi.out
    doc = Document(tep)
    o_bang = [r for t in doc.tables for h in t.rows for o in h.cells for p in o.paragraphs for r in p.runs if r.text]
    assert o_bang and {r.font.size for r in o_bang} == {Pt(M.CO_CHU_BANG_DE_CUONG)}
    than = next(p for p in doc.paragraphs if p.text.startswith("Đoạn văn thân bài"))
    assert {r.font.size for r in than.runs} == {Pt(13)}
    assert "☐" in "\n".join(_van_ban(doc))


def test_che_do_thuong_khong_dung_chuan_de_cuong(monkeypatch):
    def _khong_duoc_goi():
        raise AssertionError("chế độ thường không được nạp chuan_trinh_bay")
    monkeypatch.setattr(M, "_nap_chuan_trinh_bay", _khong_duoc_goi)
    doc = M.dung_docx(MD, 13.0)
    o_bang = [r for t in doc.tables for h in t.rows for o in h.cells for p in o.paragraphs for r in p.runs if r.text]
    assert {r.font.size for r in o_bang} == {Pt(12)}, "chế độ thường giữ bảng = cỡ thân − 1"
    assert "☐" in "\n".join(_van_ban(doc))
