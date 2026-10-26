#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tóm tắt chứng cứ CHI TIẾT bằng tiếng Việt cho từng trích dẫn có toàn văn PDF trong EndNote → Word Times New Roman.

VÌ SAO CÓ (19/10/2026, bác sĩ yêu cầu): «với mỗi trích dẫn có toàn văn với file pdf hãy bổ sung một file Word tóm tắt chi
tiết bằng tiếng Việt, font Times New Roman, theo cấu trúc chuẩn tóm tắt chứng cứ y khoa cho thực hành lâm sàng… tôi sẽ dựa
vào đó áp dụng cho thực hành». Bác sĩ DỰA VÀO bản này ⇒ máy không được để lọt một con số không có trong bài.

DÂY CHUYỀN (mỗi trích dẫn):
  1. `chuan-bi <mã>` — tách văn bản PDF trong EndNote theo TRANG, đối chiếu định danh (DOI/tiêu đề có trong PDF), tra RÚT BÀI
     bằng chuỗi 3 tầng THẬT (không trả lời bằng trí nhớ), gom ngữ cảnh dashboard (mục nào dùng bài, decision, gradeLevel). Gói
     đọc ở `~/.ebm-cache/tom-tat-chung-cu/<khoá>/` (ngoài OneDrive — chứa toàn văn, không phân phối).
  2. Một phiên/agent Claude ĐỌC TOÀN VĂN theo `HUONG_DAN_DOC` (in kèm gói) và viết `tom_tat.json` đúng khuôn `MAU`.
  2b. `kiem-cheo <mã>` — sau khi cổng tất định đạt: MỘT agent ĐỘC LẬP (ngữ cảnh mới, không soạn bản này) chấm TỪNG câu/dòng
     bảng với toàn văn (dung · suy_luan_hop_ly · ngoai_bai_co_nguon · sai · khong_co_trong_bai) → `kiem_cheo.json` gắn băm
     SHA-256 của đúng bản tom_tat.json đã chấm (bác sĩ 19/10/2026: «không được phép sai sót về chứng cứ và số liệu»).
  3. `dung <mã>` — CỔNG kiểm (fail-closed) + đòi kiểm chéo ĐỦ 100% mục, 0 sai/0 không có trong bài, rồi dựng Word:
       · mọi con số ở bảng kết quả/tác hại phải có NGUYÊN VĂN trên trang khai (±1 trang); số do máy tự tính (ARR, NNT) phải
         khai `may_tinh` và in dấu † — không lẫn với số của bài;
       · mọi trích dẫn nguyên văn ≤ 15 từ và tìm thấy trên trang khai;
       · công cụ nguy cơ sai lệch ĐÚNG thiết kế (RoB 2 chỉ cho RCT; ROBINS cho quan sát; AMSTAR-2 cho tổng quan; QUADAS cho
         chẩn đoán; AGREE II cho guideline);
       · GRADE do máy đánh giá — không bao giờ thay phân hạng của dashboard/bác sĩ; nguồn gốc này, «cần kiểm chứng» và các
         điểm «cần xác nhận» chỉ ghi ở mục Lưu ý CUỐI văn bản (bác sĩ 19/10/2026), thân văn bản không chứa các cụm đó;
       · bài đã RÚT ⇒ chặn (mã 3); định danh PDF không khớp ⇒ chặn (mã 3).
     Word lưu cạnh thư viện EndNote: `<thư mục thư viện>/TomTatChungCu/<khoá>_TomTatChungCu.docx` (đi theo OneDrive);
     JSON nguồn đã qua cổng lưu ở `TomTatChungCu/_nguon/`. Kèm bản PDF GIỮ MÀU cùng tên (`…_TomTatChungCu.pdf`, dựng bằng
     `docx_sang_pdf_giu_mau.py` — Chrome headless) vì EndNote KHÔNG xem được Word trong khung đọc (bác sĩ 21/10/2026: «chuyển
     tất cả sang file pdf để tôi có thể mở trực tiếp như file PMID pdf toàn văn»).
  4. `dung <mã> --gan-endnote --ap-dung` — THÊM bản PDF tóm tắt vào trường «File Attachments» của bản ghi, THAY liên kết Word
     cùng tên nếu có (Word vẫn nằm trong thư mục để sửa); giữ nguyên PDF toàn văn — đọc lại phải còn đủ tệp được giữ, thiếu ⇒
     ghi trả lại ngay và báo lỗi. Thiếu Chrome/Edge (không dựng được PDF) ⇒ gắn Word như trước. `pdf-hoa` làm việc này cho
     mọi bản Word đã có. Đo 19/10 trên BẢN SAO thư viện: EndNote nhận Word làm tệp LIÊN
     KẾT (đường dẫn tuyệt đối, `file_type` 2) — mở được trên máy Mac này; máy khác cần mở tệp trong thư mục OneDrive.
     Mọi lần ghi đi qua khoá ghi `khoa_ghi_endnote.py` (chờ khi bác sĩ đang thao tác EndNote). Gán `internal-pdf://` qua
     AppleScript KHÔNG được (EndNote trả «missing value») — đừng thử lại trên thư viện thật.

KHÔNG làm: không đổi `decision`/`gradeLevel` của dashboard, không sinh lời dặn bệnh nhân, không chép đoạn văn của bài (chỉ dữ
kiện + trích ≤ 15 từ), không mở/đọc nội dung DynaMed/Scopus/WoS. Mục «Gợi ý áp dụng» là GỢI Ý — áp dụng cho bệnh nhân là
Cổng A của bác sĩ.

Lệnh:
  python3 tools/tom_tat_chung_cu_word.py liet-ke [--tang 1|2|3|4] [--chua-lam] [--json]
  python3 tools/tom_tat_chung_cu_word.py mau                       # khuôn JSON + hướng dẫn đọc
  python3 tools/tom_tat_chung_cu_word.py chuan-bi <PMID|REC-n>
  python3 tools/tom_tat_chung_cu_word.py kiem <PMID|REC-n>        # chỉ chạy cổng
  python3 tools/tom_tat_chung_cu_word.py kiem-cheo <PMID|REC-n>   # dựng gói cho agent kiểm chéo độc lập
  python3 tools/tom_tat_chung_cu_word.py dung <PMID|REC-n> [--gan-endnote] [--ap-dung]
  python3 tools/tom_tat_chung_cu_word.py tien-do
  python3 tools/tom_tat_chung_cu_word.py moc --loi "<nguyên văn bác sĩ>"   # bản ghi hiện có: để nguyên, không tóm tắt
  python3 tools/tom_tat_chung_cu_word.py ngoai-le-moc --tang 1 2 --loi "<nguyên văn bác sĩ>"   # bản ghi CŨ tầng đó vẫn vào hàng chờ
  python3 tools/tom_tat_chung_cu_word.py hang-cho [--dashboard <tệp>] [--chuan-bi] [--json]   # chứng cứ MỚI chờ tóm tắt (trừ bài chỉ thuộc nhóm đề tài «NC · »)
  python3 tools/tom_tat_chung_cu_word.py pdf-hoa [--ap-dung]   # mọi Word đã có → PDF giữ màu, gắn PDF thay Word trong EndNote
Mã thoát: 0 đạt · 1 cổng có lỗi (sửa JSON) · 2 không đo được (thiếu thư viện/PDF/gói) · 3 CHẶN an toàn (rút bài/định danh).
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import importlib.util
import json
import os
import re
import sqlite3
import subprocess
import sys
import unicodedata
from urllib.parse import unquote
from pathlib import Path

GOC = Path(__file__).resolve().parents[1]
CACHE = Path.home() / ".ebm-cache" / "tom-tat-chung-cu"
TEN_THU_MUC_RA = "TomTatChungCu"
PHIEN_BAN_KHUON = 2  # v2 (19/10/2026): thêm `diem_nhan`; bỏ đối chiếu dashboard · trích nguyên văn · câu hỏi mở
TOI_DA_TU_TRICH = 15
LOAI_THIET_KE = ("rct", "sr_ma", "cohort", "case_control", "cross_sectional", "diagnostic", "guideline", "nhan_thuoc", "khac")
TEN_THIET_KE = {"rct": "Thử nghiệm lâm sàng ngẫu nhiên có đối chứng", "sr_ma": "Tổng quan hệ thống / phân tích gộp",
                "cohort": "Nghiên cứu thuần tập", "case_control": "Nghiên cứu bệnh–chứng", "cross_sectional": "Nghiên cứu cắt ngang",
                "diagnostic": "Nghiên cứu độ chính xác chẩn đoán", "guideline": "Hướng dẫn thực hành lâm sàng (guideline)",
                "nhan_thuoc": "Nhãn thuốc / thông tin kê đơn", "khac": "Thiết kế khác"}
# Công cụ nguy cơ sai lệch HỢP LỆ theo thiết kế (CLAUDE.md §6.3: RoB 2 chỉ cho RCT). So khớp không phân biệt hoa/thường,
# bỏ khoảng trắng/gạch — «ROBINS-I V2» khớp «robinsi».
CONG_CU_HOP_LE = {
    "rct": ("rob2",),
    "sr_ma": ("amstar2", "robis"),
    "cohort": ("robinsi", "robinse", "newcastleottawa", "nos"),
    "case_control": ("robinse", "robinsi", "newcastleottawa", "nos"),
    "cross_sectional": ("robinse", "jbi", "axis", "newcastleottawa"),
    "diagnostic": ("quadas2", "quadas3", "quadasc"),
    "guideline": ("agreeii", "agree2"),
    "nhan_thuoc": ("khongapdung",),
    "khac": (),
}
MUC_GRADE = ("cao", "trung bình", "thấp", "rất thấp", "không áp dụng")
TRUONG_SO = ("nhom_can_thiep", "nhom_so_sanh", "hieu_ung", "ktc95", "p", "tuyet_doi", "nnt")
# Ô «ĐIỂM NHẤN» mở đầu mỗi mục của Word (bác sĩ 19/10/2026). `tac_hai` không có ở guideline/nhãn thuốc.
KHOA_DIEM_NHAN = ("thong_tin", "pico", "phuong_phap", "ket_qua", "tac_hai", "nguy_co_sai_lech", "grade", "ap_dung_vn", "goi_y",
                  "manh_han_che")
TOI_DA_TU_DIEM_NHAN = 70
DANH_GIA_TAC_HAI = ("bat_loi", "loi", "tuong_duong")
MIEN_TRU = "Bản tóm tắt do MÁY (Claude) soạn từ toàn văn — CẦN BÁC SĨ KIỂM CHỨNG trước khi áp dụng cho người bệnh."

MAU = {
    "phien_ban_khuon": PHIEN_BAN_KHUON,
    "khoa": "PMID-10471456",
    "pmid": "10471456",
    "doi": "10.1056/NEJM199909023411001",
    "loai_thiet_ke": "rct",
    "thiet_ke_mo_ta": "Thử nghiệm ngẫu nhiên, mù đôi, đa trung tâm, đối chứng giả dược",
    "linh_vuc": "Tim mạch — suy tim",
    "cau_hoi_lam_sang": {"loai": "điều trị", "P": "…", "I": "…", "C": "…", "O": "…"},
    "ket_luan_thuc_hanh": ["Mỗi gạch một kết luận thực hành, giữ mệnh đề điều kiện của bài (quần thể, ngưỡng, chống chỉ định)."],
    "boi_canh": {"dia_diem": "…", "thoi_gian": "…", "dang_ky": "NCT… hoặc «không khai»", "tai_tro": "…",
                 "xung_dot_loi_ich": "…"},
    "doi_tuong": {"tieu_chuan_chon": ["…"], "tieu_chuan_loai": ["…"], "co_mau": "…", "dac_diem_nen": "…"},
    "can_thiep_so_sanh": {"can_thiep": "thuốc · liều · đường dùng · thời gian", "so_sanh": "…", "dong_tri_lieu": "…"},
    "ket_cuc": {"chinh": "…", "phu": ["…"], "theo_doi": "…"},
    "phan_tich": "ITT/PP, mô hình, xử lý dữ liệu thiếu, dừng sớm…",
    "ket_qua": [{"ket_cuc": "Tử vong mọi nguyên nhân", "chinh": True, "nhom_can_thiep": "284/822 (35%)",
                 "nhom_so_sanh": "386/841 (46%)", "hieu_ung": "RR 0.70", "ktc95": "0.60–0.82", "p": "<0.001",
                 "tuyet_doi": "ARR 11%", "nnt": "9", "may_tinh": ["tuyet_doi", "nnt"], "trang": 3}],
    "tac_hai": [{"ket_cuc": "Vú to/đau vú ở nam", "chinh": False, "nhom_can_thiep": "10%", "nhom_so_sanh": "1%",
                 "hieu_ung": "", "ktc95": "", "p": "<0.001", "tuyet_doi": "", "nnt": "", "may_tinh": [], "trang": 4,
                 "danh_gia": "bat_loi"}],
    "tac_hai_ghi_chu": "Ghi khi bài không báo cáo tác hại có cấu trúc.",
    "phan_nhom": ["Phân nhóm (kèm kiểm định tương tác nếu bài báo cáo)."],
    "khuyen_cao": [],
    "nguy_co_sai_lech": {"cong_cu": "RoB 2", "mien": [{"ten": "Quy trình ngẫu nhiên hoá", "danh_gia": "thấp",
                                                         "ly_do": "…", "trang": 2}], "tong_the": "nguy cơ thấp"},
    "grade_de_xuat": [{"ket_cuc": "Tử vong mọi nguyên nhân", "muc": "cao", "ly_do": "lý do hạ/nâng từng miền GRADE"}],
    "kha_nang_ap_dung": {"tuong_dong": "…", "khac_biet": "…", "trien_khai_viet_nam": "thuốc sẵn có, BHYT, xét nghiệm theo dõi…"},
    "goi_y_ap_dung": {"nen_lam": ["…"], "khong_nen": ["…"], "dieu_kien": ["…"], "theo_doi": ["…"]},
    "diem_manh": ["…"],
    "han_che": ["…"],
    "diem_nhan": {"thong_tin": "…", "pico": "…", "phuong_phap": "…", "ket_qua": "…", "tac_hai": "…",
                  "nguy_co_sai_lech": "…", "grade": "…", "ap_dung_vn": "…", "goi_y": "…", "manh_han_che": "…"},
    "can_xac_nhan": ["Điểm chưa chắc (chi phí, BHYT, chứng cứ khác cần tra) — chỉ in ở mục Lưu ý cuối văn bản"],
    "nguoi_soan": "Claude (máy) — đề xuất, chưa qua bác sĩ",
}
# Khoá bắt buộc theo thiết kế (guideline/nhãn thuốc không có nhóm can thiệp–so sánh theo nghĩa thử nghiệm).
_BAT_BUOC_CHUNG = ("khoa", "loai_thiet_ke", "thiet_ke_mo_ta", "cau_hoi_lam_sang", "ket_luan_thuc_hanh", "boi_canh",
                   "nguy_co_sai_lech", "grade_de_xuat", "kha_nang_ap_dung", "goi_y_ap_dung", "diem_manh", "han_che",
                   "diem_nhan")
_BAT_BUOC_NGHIEN_CUU = ("doi_tuong", "ket_cuc", "phan_tich", "ket_qua")
_BAT_BUOC_THEO_THIET_KE = {
    "rct": _BAT_BUOC_NGHIEN_CUU + ("can_thiep_so_sanh",),
    "sr_ma": _BAT_BUOC_NGHIEN_CUU + ("can_thiep_so_sanh",),
    "cohort": _BAT_BUOC_NGHIEN_CUU,
    "case_control": _BAT_BUOC_NGHIEN_CUU,
    "cross_sectional": _BAT_BUOC_NGHIEN_CUU,
    "diagnostic": _BAT_BUOC_NGHIEN_CUU,
    "guideline": ("khuyen_cao",),
    "nhan_thuoc": ("khuyen_cao",),
    "khac": ("ket_qua",),
}

HUONG_DAN_DOC = """# Hướng dẫn ĐỌC TOÀN VĂN và viết `tom_tat.json` — khuôn phiên bản {pb}

Bác sĩ sẽ DỰA VÀO bản này để thực hành. Đọc HẾT văn bản trong `van_ban.txt` (mỗi trang mở đầu bằng `=== TRANG n ===`);
guideline dài: đọc phần phương pháp + MỌI khuyến cáo liên quan câu hỏi của dashboard trong `ngu_canh.json`.

Luật (cổng `dung` kiểm máy — sai là KHÔNG ra Word):
1. MỌI con số trong `ket_qua` và `tac_hai` (các trường `nhom_can_thiep`, `nhom_so_sanh`, `hieu_ung`, `ktc95`, `p`, `tuyet_doi`,
   `nnt`) phải chép ĐÚNG như trong bài, và có trên trang `trang` (±1). Dấu thập phân giữ như bài (0.70 / 0·70 đều được).
   Số do bạn TỰ TÍNH (ARR, NNT, NNH…) ⇒ ghi tên trường vào `may_tinh` (Word in dấu †) và ghi cách tính ở `phan_tich`.
   Bài không báo cáo ⇒ để chuỗi rỗng, KHÔNG suy đoán.
2. Không chép đoạn văn của bài (bản quyền) — chỉ dữ kiện bằng lời tiếng Việt. (`trich_nguyen_van` là tuỳ chọn, không in ra
   Word; nếu ghi thì tối đa {tu} từ mỗi câu, đúng chữ trên trang khai.)
3. `nguy_co_sai_lech.cong_cu` đúng thiết kế: RCT → «RoB 2»; tổng quan/gộp → «AMSTAR-2» (hoặc ROBIS); thuần tập/bệnh–chứng →
   «ROBINS-I» / «ROBINS-E» (hoặc Newcastle-Ottawa); chẩn đoán → «QUADAS-2»/«QUADAS-3»; guideline → «AGREE II» (ghi rõ Miền 3
   Rigour nếu guideline không mô tả); nhãn thuốc → «Không áp dụng». Mỗi miền: `danh_gia` + `ly_do` + `trang`.
4. `grade_de_xuat`: mức ∈ {{cao, trung bình, thấp, rất thấp, không áp dụng}} cho TỪNG kết cục chính, lý do theo 5 miền hạ
   (nguy cơ sai lệch · không nhất quán · gián tiếp · không chính xác · sai lệch xuất bản) và miền nâng. Nguồn gốc «máy đánh giá»
   được ghi MỘT lần ở mục Lưu ý cuối văn bản — thân văn bản không viết «đề xuất của máy». Guideline thì ghi phân hạng NGUYÊN
   BẢN của guideline trong `khuyen_cao`, không tự gán.
5. `ket_luan_thuc_hanh` (2–8 gạch) và `goi_y_ap_dung`: giữ mệnh đề điều kiện của bài (quần thể, ngưỡng, chống chỉ định,
   «ngoài thai kỳ»…); số liệu gắn đúng kết cục nó đo; không khẳng định vượt bài. Không viết lời dặn bệnh nhân, không bịa liều —
   liều chỉ chép từ bài kèm trang.
6. `kha_nang_ap_dung.trien_khai_viet_nam`: sẵn có thuốc/xét nghiệm, BHYT, khác biệt quần thể — không đoán. Điều CHƯA CHẮC
   (chi phí, BHYT, sẵn có, chứng cứ khác cần tra) KHÔNG viết vào thân văn bản: ghi từng điểm vào `can_xac_nhan` (in ở mục Lưu
   ý cuối văn bản). Bác sĩ 19/10/2026: thân văn bản không chứa «cần kiểm chứng», «cần xác nhận», «[CẦN…]», «đề xuất của
   máy» — cổng chặn.
7. `diem_nhan` — ô ĐIỂM NHẤN mở đầu từng mục của Word, mỗi ô 1–2 câu (≤ {tdn} từ) nói điều QUAN TRỌNG NHẤT của mục:
   {kdn} (guideline/nhãn thuốc bỏ `tac_hai`). Nhấn trong MỌI trường chữ bằng: `**đậm**` ·
   `[!cảnh báo!]` (đỏ đậm — chống chỉ định, tác hại, nhóm bị loại) · `[+lợi ích+]` (xanh lá đậm — lợi ích chính). Số trong `ket_luan_thuc_hanh`, `diem_nhan`, `goi_y_ap_dung` cũng phải có NGUYÊN VĂN trong bài
   (hoặc trong bảng kết quả/tác hại) — viết đúng số của bài (93.1%, không «≈ 93%»; không tự cộng tổng). Luật này áp cho
   MỌI trường chữ, không chỉ phần tóm lược. Số lấy từ nguồn KHÁC (nhãn thuốc,
   guideline, dashboard) phải nằm trong cùng một gạch có nhãn `[ngoài bài: <nguồn/PMID>]` — Word in xám nghiêng.
   Mỗi dòng `tac_hai` thêm `danh_gia` ∈ {{bat_loi, loi, tuong_duong}} (can thiệp nhiều tác hại hơn · ít hơn · không khác) để tô màu.
   Đọc `ngu_canh.json` để biết bài đang được dashboard dùng thế nào (decision, gradeLevel) — chỉ để hiểu bối cảnh.
8. guideline/nhãn thuốc: điền `khuyen_cao` (mỗi khuyến cáo liên quan: nội dung tiếng Việt, `luc_khuyen_cao`, `muc_chung_cu`
   NGUYÊN BẢN, `trang`, `lien_quan_dashboard`); `ket_qua` để [].
9. Không PII. Không thông tin ngoài bài trừ khi ghi nguồn (PMID/DOI) ngay trong câu. Không viết so sánh nhất («lớn nhất»,
   «đầu tiên») hay nhận định về thực hành ở Việt Nam nếu bài không viết — gắn `[ngoài bài: nguồn]` khi lấy từ nơi khác.
   Không dùng «≈», «~» — ghi đúng số của bài, báo cáo theo nhóm thì ghi đủ HAI nhóm (93.1% / 92.9%), không tự cộng/làm tròn.
   Sau bạn, một agent ĐỘC LẬP sẽ chấm từng câu với toàn văn; một câu «sai» hay «không có trong bài» là Word không được dựng.
10. Viết tiếng Việt chuẩn y khoa, thuật ngữ chưa có tương đương thì tiếng Việt kèm tiếng Anh trong ngoặc lần đầu.

Khuôn JSON (thay «…» bằng nội dung thật; giữ đủ khoá):
"""


def huong_dan_doc() -> str:
    return HUONG_DAN_DOC.format(pb=PHIEN_BAN_KHUON, tu=TOI_DA_TU_TRICH, tdn=TOI_DA_TU_DIEM_NHAN,
                                kdn=" · ".join(f"`{k}`" for k in KHOA_DIEM_NHAN))


# ───────────────────────── tiện ích chung ─────────────────────────

def _nap(ten: str, tep: str):
    sp = importlib.util.spec_from_file_location(ten, Path(__file__).resolve().parent / tep)
    mod = importlib.util.module_from_spec(sp)
    sp.loader.exec_module(mod)
    return mod


_MOD: dict = {}


def mod(ten: str):
    tep = {"sk": "endnote_suc_khoe.py", "nec": "nap_endnote_chu_de.py", "tt": "tom_tat_endnote_vi.py",
           "sx": "so_xac_minh_nguon.py"}[ten]
    if ten not in _MOD:
        _MOD[ten] = _nap(f"_ttcc_{ten}", tep)
    return _MOD[ten]


def khoa_cua(pmid: str | None, rec: int) -> str:
    return f"PMID-{pmid}" if pmid and re.fullmatch(r"\d{1,9}", str(pmid)) else f"REC-{rec}"


def thu_muc_ra(thu_vien: Path) -> Path:
    return thu_vien.parent / TEN_THU_MUC_RA


def tep_word(thu_vien: Path, khoa: str) -> Path:
    return thu_muc_ra(thu_vien) / f"{khoa}_TomTatChungCu.docx"


def tep_pdf(tep_docx: Path) -> Path:
    """Bản PDF đứng cạnh bản Word (cùng tên, đuôi .pdf) — EndNote mở PDF ngay trong khung đọc, Word thì không."""
    return tep_docx.with_suffix(".pdf")


def pdf_con_moi(tep_docx: Path) -> bool:
    """PDF đã có và KHÔNG cũ hơn Word (Word dựng lại sau ⇒ phải dựng lại PDF, không gắn bản lệch)."""
    pdf = tep_pdf(tep_docx)
    return pdf.is_file() and pdf.stat().st_mtime >= tep_docx.stat().st_mtime


def dung_pdf(tep_docx: Path, *, chay=subprocess.run) -> Path | None:
    """Dựng PDF GIỮ MÀU từ bản Word bằng `docx_sang_pdf_giu_mau.py` (Chrome/Edge headless, cùng công cụ bước ⑤ của bộ
    năm). Thiếu trình duyệt/lỗi/PDF ra cũ hơn Word ⇒ None (người gọi gắn Word thay vì gắn PDF lệch nội dung)."""
    ra = tep_pdf(tep_docx)
    moi_truong = dict(os.environ)
    pandoc_rieng = Path.home() / ".local" / "bin"  # pandoc của máy Mac nằm ở đây (bộ nhớ toolchain docx)
    if pandoc_rieng.is_dir():
        moi_truong["PATH"] = str(pandoc_rieng) + os.pathsep + moi_truong.get("PATH", "")
    try:
        kq = chay([sys.executable, str(GOC / "tools" / "docx_sang_pdf_giu_mau.py"), str(tep_docx), "--ra", str(ra)],
                  capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=300, env=moi_truong)
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"⚪ không dựng được PDF cho {tep_docx.name}: {e}")
        return None
    if kq.returncode != 0 or not pdf_con_moi(tep_docx):
        print(f"⚪ không dựng được PDF cho {tep_docx.name} (mã {kq.returncode}): {(kq.stdout or '')[-200:].strip()}")
        return None
    return ra


def _chuan_chu(s: str) -> str:
    """Chuẩn hoá để so khớp chữ trong PDF: NFKC, bỏ gạch nối cuối dòng, gộp khoảng trắng, chữ thường, thống nhất nháy/gạch."""
    s = unicodedata.normalize("NFKC", s or "")
    s = s.replace("­", "")
    s = re.sub(r"-\s*\n\s*", "", s)
    s = s.translate(str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
                                   "−": "-", "·": ".", "⋅": "."}))
    return re.sub(r"\s+", " ", s).strip().lower()


def _chuan_so(s: str) -> str:
    """Chuẩn hoá cách viết số hai phía như nhau: «0·70»→0.70; phân cách nghìn «10,000» hoặc «10 000» (dấu cách MẢNH/không
    ngắt) → 10000; phẩy thập phân → chấm. KHÔNG coi dấu cách THƯỜNG là phân cách nghìn: văn bản tách từ bảng PDF đặt các số
    cạnh nhau bằng dấu cách («122±15 121±15», «LCZ696 100 mg») — gộp chúng tạo ra số không có thật (đo 19/10: «15121»,
    «696100»). Bỏ phân cách nghìn TRƯỚC khi NFKC (NFKC đổi dấu cách mảnh/không ngắt thành dấu cách thường)."""
    s = re.sub(r"(?<=\d)[,\u2009\u202f\u00a0](?=\d{3}(?!\d))", "", s or "")
    s = unicodedata.normalize("NFKC", s)
    s = re.sub(r"(?<=\d)[\u00b7\u22c5](?=\d)", ".", s)
    return re.sub(r"(?<=\d),(?=\d)", ".", s)


_SO = re.compile(r"\d+(?:\.\d+)?")


def cac_so(s: str) -> list[str]:
    """Các số trong chuỗi; số nguyên bỏ số 0 đứng đầu («08/12/2009» ≡ «December 8, 2009»)."""
    return [x if "." in x else (x.lstrip("0") or "0") for x in _SO.findall(_chuan_so(s))]


def _so_khop(so: str, tap: set[str]) -> bool:
    if so in tap:
        return True
    if "." in so:  # 0.6 ≡ 0.60 — chỉ nới cho số thập phân, số nguyên phải trùng nguyên văn
        try:
            f = float(so)
        except ValueError:
            return False
        return any("." in t and abs(float(t) - f) < 1e-12 for t in tap)
    return False


# ───────────────────────── văn bản PDF ─────────────────────────

def tach_trang(pdf: Path) -> list[str]:
    """Văn bản từng trang (pypdf). Trang không trích được ⇒ chuỗi rỗng (PDF ảnh)."""
    import logging  # noqa: PLC0415

    import pypdf  # noqa: PLC0415
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    doc = pypdf.PdfReader(str(pdf))
    ra = []
    for tr in doc.pages:
        try:
            ra.append(tr.extract_text() or "")
        except Exception:  # noqa: BLE001 — một trang hỏng không làm hỏng cả bài
            ra.append("")
    return ra


def ghep_van_ban(trang: list[str]) -> str:
    return "\n".join(f"=== TRANG {i} ===\n{t}" for i, t in enumerate(trang, 1))


def doc_van_ban(goi: Path) -> list[str]:
    """Đọc lại `van_ban.txt` thành danh sách trang (trang 1 ở chỉ số 0)."""
    tho = (goi / "van_ban.txt").read_text(encoding="utf-8")
    phan = re.split(r"^=== TRANG (\d+) ===\n", tho, flags=re.M)
    trang: dict[int, str] = {}
    for i in range(1, len(phan) - 1, 2):
        trang[int(phan[i])] = phan[i + 1]
    return [trang.get(i, "") for i in range(1, max(trang, default=0) + 1)]


def khop_dinh_danh(trang: list[str], doi: str | None, tieu_de: str | None) -> tuple[bool, str]:
    """PDF có đúng là bài này không — DOI xuất hiện trong văn bản, hoặc ≥ 60% từ (≥ 4 chữ) của tiêu đề có ở 3 trang đầu."""
    dau = _chuan_chu(" ".join(trang[:3]))
    toan = _chuan_chu(" ".join(trang)).replace(" ", "")
    if doi and _chuan_chu(doi).replace(" ", "") in toan:
        return True, "DOI có trong PDF"
    tu = [w for w in re.findall(r"[a-z0-9]+", _chuan_chu(tieu_de or "")) if len(w) >= 4]
    if tu:
        ty_le = sum(w in dau for w in tu) / len(tu)
        if ty_le >= 0.6:
            return True, f"{ty_le:.0%} từ của tiêu đề có ở 3 trang đầu"
        return False, f"chỉ {ty_le:.0%} từ của tiêu đề có ở 3 trang đầu và không thấy DOI"
    return False, "bản ghi không có DOI/tiêu đề để đối chiếu"


# ───────────────────────── EndNote (chỉ đọc) ─────────────────────────

def ban_ghi_co_pdf(thu_vien: Path) -> list[dict]:
    """Mọi bản ghi NGOÀI thùng rác có ít nhất một tệp PDF — đọc `sdb.eni` chế độ chỉ-đọc."""
    sk = mod("sk")
    c = sk._ket_noi(thu_vien)
    try:
        cot = {r[1] for r in c.execute("pragma table_info(refs)")}
        tac_gia = "author" if "author" in cot else "''"
        tap_chi = "secondary_title" if "secondary_title" in cot else "''"
        hang = c.execute(
            f"select r.id, r.accession_number, r.electronic_resource_number, r.title, r.year, {tac_gia}, {tap_chi}, "
            "f.file_path from refs r join file_res f on f.refs_id = r.id where r.trash_state = 0 "
            "and lower(f.file_path) like '%.pdf' order by r.id, f.file_pos").fetchall()
    finally:
        c.close()
    ra: dict[int, dict] = {}
    pdf_goc = thu_vien.with_suffix(".Data") / "PDF"
    for rec, acc, erd, tieu_de, nam, tg, tc, fp in hang:
        pmid = acc if acc and re.fullmatch(r"\d{1,9}", str(acc).strip()) else None
        doi = mod("nec").chuan_doi(erd) if erd else None
        m = ra.setdefault(rec, {"rec": rec, "pmid": pmid, "doi": doi, "tieu_de": (tieu_de or "").strip(),
                                "nam": str(nam or ""), "tac_gia": (tg or "").split("\r")[0].strip(),
                                "tap_chi": (tc or "").strip(), "pdf": []})
        m["pdf"].append(str(pdf_goc / fp))
    for m in ra.values():
        m["khoa"] = khoa_cua(m["pmid"], m["rec"])
    return list(ra.values())


def tim_ban_ghi(thu_vien: Path, ma: str) -> dict | None:
    ma = ma.strip()
    so = re.sub(r"^(PMID[-:]?|REC[-:]?)", "", ma, flags=re.I)
    la_rec = bool(re.match(r"^REC", ma, flags=re.I))
    for m in ban_ghi_co_pdf(thu_vien):
        if (la_rec and str(m["rec"]) == so) or (not la_rec and m["pmid"] == so):
            return m
    return None


# ───────────────────────── ngữ cảnh dashboard + ưu tiên ─────────────────────────

def goc_du_lieu(ten: str) -> Path:
    """Gốc dữ liệu ngoài git (lồng rồi anh em) qua `ban_sao_tran.duong_goc` — cùng cách định vị với `so_toan_van.py`."""
    try:
        return _nap("_ttcc_bst", "ban_sao_tran.py").duong_goc(ten, GOC) or (GOC / ten)
    except Exception:  # noqa: BLE001 — thiếu bộ định vị ⇒ vị trí lồng
        return GOC / ten


def ngu_canh_dashboard() -> dict[str, dict]:
    files = sorted(goc_du_lieu("EBM-Dashboards").glob("WebDashboard_*.html"))
    if not files:
        return {}
    return mod("tt").ngu_canh(files, mod("nec")._nap_doc_data())


def _ngu_canh_cua(ds: dict[str, dict], pmid: str | None, doi: str | None) -> list[dict]:
    v = ds.get(f"pmid:{pmid}") if pmid else None
    if v is None and doi:
        v = ds.get(f"doi:{doi}") or next((x for x in ds.values() if x.get("doi") == doi), None)
    ra = []
    for nc in (v or {}).get("ngu_canh", []):
        it = nc.get("item") or {}
        ra.append({"dashboard": nc.get("dashboard"), "ngay": nc.get("ngay"), "chu_de": nc.get("chu_de"),
                   "muc": nc.get("muc"), "tieu_de_muc": nc.get("tieu_de"), "vai": nc.get("vai"),
                   "ban_cu": bool(nc.get("ban_cu")), "decision": it.get("decision"), "gradeLevel": it.get("gradeLevel"),
                   "gradeSource": it.get("gradeSource"), "design": it.get("design"), "population": it.get("population"),
                   "action": it.get("action"), "monitoring": it.get("monitoring"), "vn": it.get("vn"),
                   "tham_khao": nc.get("tham_khao")})
    return ra


def tang_uu_tien(ngu_canh: list[dict]) -> int:
    """1 = nguồn chính của mục apply · 2 = consider · 3 = notyet/khác · 4 = chỉ tham khảo hoặc không có dashboard."""
    moi = [x for x in ngu_canh if not x.get("ban_cu")] or ngu_canh
    dec = {str(x.get("decision") or "") for x in moi if x.get("vai") == "muc"}
    if "apply" in dec:
        return 1
    if "consider" in dec:
        return 2
    if any(x.get("vai") == "muc" for x in moi):
        return 3
    return 4


# ───────────────────────── rút bài ─────────────────────────

def kiem_rut_bai(pmid: str | None, doi: str | None) -> dict:
    """Tra bằng CÔNG CỤ (chuỗi 3 tầng của repo y khoa; DOI-only ⇒ nền Retraction Watch ngoại tuyến). Không tra được ⇒
    trạng thái «chua_kiem» — tuyệt đối không coi là «chưa bị rút»."""
    luc = dt.datetime.now().isoformat(timespec="seconds")
    try:
        sx = mod("sx")
        if pmid:
            kq = (sx.kiem_rut_bai([pmid]) or {}).get(pmid)
            if isinstance(kq, dict) and kq.get("status") in ("ok", "retracted"):
                nguon = kq.get("source") or ",".join(kq.get("sources_tried") or [])
                if kq["status"] == "retracted" and kq.get("withdrawn_correction_notice"):
                    # 24/10/2026: thông báo rút là BẢN ĐÍNH CHÍNH bị rút mà bác sĩ ĐÃ KÝ xem xét (rut-bai-da-xem-xet.json, đúng
                    # vân tay tập thông báo) — cùng hàm của sổ xác minh/cổng verify_dashboard. Đo: guideline CCS/CHFS 2025
                    # (PMID 41110921, bác sĩ ký 24/09) bị `chuan-bi` chặn «ĐÃ BỊ RÚT» trong khi cổng dashboard cho qua.
                    # Chưa ký / vân tay lệch / không đối chiếu được ⇒ vẫn «da_rut» (fail-closed).
                    tam: dict = {}
                    sx._gan_dau_hieu_thong_bao(tam, kq)
                    if tam.get("sua_loi_bi_rut") and sx._dinh_chinh_da_ky(f"pmid:{pmid}", tam):
                        return {"trang_thai": "dinh_chinh_bi_rut_da_ky", "nguon": nguon, "luc": luc,
                                "thong_bao_ids": tam.get("thong_bao_ids") or []}
                return {"trang_thai": "da_rut" if kq["status"] == "retracted" else "khong_thay_rut",
                        "nguon": nguon, "luc": luc}
        elif doi:
            kq = sx.rut_bai_retraction_watch_ngoai_tuyen_doi([doi])
            if kq is not None:
                return {"trang_thai": "da_rut" if kq else "khong_thay_rut_rw", "nguon": "retraction_watch (ngoại tuyến)",
                        "luc": luc}
    except Exception as e:  # noqa: BLE001 — mọi lỗi tra ⇒ KHÔNG BIẾT
        return {"trang_thai": "chua_kiem", "nguon": f"lỗi tra: {str(e)[:160]}", "luc": luc}
    return {"trang_thai": "chua_kiem", "nguon": "công cụ không trả kết luận", "luc": luc}


NHAN_RUT_BAI = {"khong_thay_rut": "Không thấy thông báo rút bài (chuỗi 3 tầng: Retraction Watch · PubMed · Europe PMC)",
                "khong_thay_rut_rw": "Không thấy trong Retraction Watch (chỉ tra ngoại tuyến theo DOI — bài không có PMID)",
                "da_rut": "ĐÃ BỊ RÚT — KHÔNG ĐƯỢC DÙNG",
                "dinh_chinh_bi_rut_da_ky": "Cờ rút bài là BẢN ĐÍNH CHÍNH bị rút — bác sĩ đã ký xem xét (rut-bai-da-xem-xet.json, "
                                           "đúng vân tay thông báo); bài gốc không bị rút",
                "chua_kiem": "CHƯA KIỂM ĐƯỢC RÚT BÀI — tra lại trước khi dùng"}


# ───────────────────────── chuẩn bị gói đọc ─────────────────────────

def chuan_bi(thu_vien: Path, ma: str, *, ngu_canh_tat_ca: dict | None = None, rut_bai=kiem_rut_bai) -> tuple[int, Path | None]:
    bg = tim_ban_ghi(thu_vien, ma)
    if bg is None:
        print(f"⚪ không thấy bản ghi có PDF cho «{ma}» trong thư viện {thu_vien.name}")
        return 2, None
    goi = CACHE / bg["khoa"]
    goi.mkdir(parents=True, exist_ok=True)
    pdf = Path(bg["pdf"][0])
    if not pdf.is_file():
        print(f"⚪ tệp PDF vắng trên đĩa: {pdf} (OneDrive chưa tải về?)")
        return 2, None
    trang = tach_trang(pdf)
    so_ky_tu = sum(len(t) for t in trang)
    (goi / "van_ban.txt").write_text(ghep_van_ban(trang), encoding="utf-8")
    khop, ly_do = khop_dinh_danh(trang, bg["doi"], bg["tieu_de"])
    ngu_canh = _ngu_canh_cua(ngu_canh_tat_ca if ngu_canh_tat_ca is not None else ngu_canh_dashboard(), bg["pmid"], bg["doi"])
    rb = rut_bai(bg["pmid"], bg["doi"])
    nc = {"ban_ghi": {k: bg[k] for k in ("rec", "pmid", "doi", "tieu_de", "nam", "tac_gia", "tap_chi", "khoa")},
          "pdf": {"tep": pdf.name, "sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(), "so_trang": len(trang),
                  "so_ky_tu": so_ky_tu, "trang_trong": [i for i, t in enumerate(trang, 1) if len(t.strip()) < 40]},
          "dinh_danh": {"khop": khop, "ly_do": ly_do}, "rut_bai": rb, "tang_uu_tien": tang_uu_tien(ngu_canh),
          "ngu_canh_dashboard": ngu_canh, "chuan_bi_luc": dt.datetime.now().isoformat(timespec="seconds")}
    (goi / "ngu_canh.json").write_text(json.dumps(nc, ensure_ascii=False, indent=2), encoding="utf-8")
    (goi / "HUONG-DAN-DOC.md").write_text(
        huong_dan_doc() + "\n```json\n"
        + json.dumps(MAU, ensure_ascii=False, indent=2) + "\n```\n", encoding="utf-8")
    print(f"Gói đọc: {goi}")
    print(f"  {bg['khoa']} · recNum {bg['rec']} · {len(trang)} trang · {so_ky_tu} ký tự · tầng ưu tiên {nc['tang_uu_tien']}")
    print(f"  định danh: {'✓' if khop else '✗'} {ly_do} · rút bài: {NHAN_RUT_BAI.get(rb['trang_thai'], rb['trang_thai'])}")
    if so_ky_tu < 2000:
        print("  ⚠ PDF gần như không có lớp chữ (PDF ảnh) — không tóm tắt máy được; bác sĩ đọc trực tiếp")
        return 2, goi
    if rb["trang_thai"] == "da_rut":
        print("  ⛔ BÀI ĐÃ BỊ RÚT — không soạn tóm tắt để áp dụng")
        return 3, goi
    if not khop:
        print("  ⛔ PDF có thể KHÔNG phải bài này — kiểm tệp đính kèm trước khi tóm tắt")
        return 3, goi
    print(f"  Đọc: {goi / 'HUONG-DAN-DOC.md'} → viết {goi / 'tom_tat.json'} → `dung {bg['khoa']}`")
    return 0, goi


# ───────────────────────── kiểm TẤT ĐỊNH nhất quán số liệu một dòng kết quả ─────────────────────────
# Bác sĩ 19/10/2026: «không được phép sai sót về chứng cứ và số liệu». Số có trong bài vẫn có thể bị GẮN NHẦM (đảo hai nhóm,
# lấy KTC của dòng khác). Các phép dưới bắt mâu thuẫn nội tại của một dòng — không cần hiểu ngữ nghĩa.

_SO_CO_DAU = r"[-−–]?\d+(?:[.,·]\d+)?"
_KTC = re.compile(rf"({_SO_CO_DAU})\s*(?:–|—|-|to|đến|,)\s*({_SO_CO_DAU})")
_TY_SO = re.compile(r"\b(HR|RR|OR|IRR|aHR|aOR|aRR|SHR|rate ratio|hazard ratio|odds ratio|risk ratio|tỷ số)\b", re.I)
SO_GAN_DUNG = re.compile(r"[≈~∼]\s*\d|\bkhoảng\s+\d+(?:[.,]\d+)?\s*%")


def _f(s: str) -> float | None:
    try:
        return float(_chuan_so(s.replace("−", "-").replace("–", "-")))
    except (TypeError, ValueError):
        return None


def _ty_le_nhom(s: str) -> float | None:
    """Tỷ lệ % của một nhóm: «914 (21.8%)», «914 (21.8)», «10.7%»."""
    s = _chuan_so(str(s or ""))
    m = re.search(r"(\d+(?:\.\d+)?)\s*%", s) or re.search(r"\((\d+(?:\.\d+)?)\)", s)
    return float(m.group(1)) if m else None


def _gia_tri_p(s: str) -> tuple[str, float] | None:
    s = _chuan_so(str(s or "")).replace(" ", "")
    m = re.search(r"([<≤>=]?)(\d*\.?\d+)", s)
    return (m.group(1), float(m.group(2))) if m else None


def kiem_nhat_quan_dong(h: dict, ten: str) -> list[str]:
    loi: list[str] = []
    hu = str(h.get("hieu_ung") or "")
    la_ty_so = bool(_TY_SO.search(hu))
    so_hu = re.findall(_SO_CO_DAU, hu)
    uoc = _f(so_hu[0]) if so_hu else None
    m = _KTC.search(str(h.get("ktc95") or ""))
    lo, hi = (_f(m.group(1)), _f(m.group(2))) if m else (None, None)
    if uoc is not None and lo is not None and hi is not None:
        if lo > hi:
            lo, hi = hi, lo
        if not lo - 1e-9 <= uoc <= hi + 1e-9:
            loi.append(f"{ten}: ước lượng {uoc} nằm NGOÀI KTC 95% {lo}–{hi} — KTC của dòng khác, hay chép nhầm?")
        moc = 1.0 if la_ty_so else 0.0
        loai_tru_moc = hi < moc or lo > moc
        p = _gia_tri_p(h.get("p"))
        if p is not None:
            dau, gt = p
            if loai_tru_moc and gt > 0.10 and dau not in ("<", "≤"):
                loi.append(f"{ten}: KTC 95% {lo}–{hi} không chứa {moc} nhưng p = {gt} — mâu thuẫn, kiểm lại")
            if not loai_tru_moc and (gt < 0.01 or (dau in ("<", "≤") and gt <= 0.01)):
                loi.append(f"{ten}: KTC 95% {lo}–{hi} chứa {moc} nhưng p {dau}{gt} — mâu thuẫn, kiểm lại")
    pi, pc = _ty_le_nhom(h.get("nhom_can_thiep")), _ty_le_nhom(h.get("nhom_so_sanh"))
    if la_ty_so and uoc is not None and pi is not None and pc is not None and abs(pi - pc) >= 0.5:
        if (pi < pc and uoc > 1) or (pi > pc and uoc < 1):
            loi.append(f"{ten}: tỷ lệ nhóm can thiệp {pi}% so với nhóm so sánh {pc}% ngược chiều {hu} — đảo hai nhóm?")
    dg = h.get("danh_gia")
    if dg in ("bat_loi", "loi") and pi is not None and pc is not None and abs(pi - pc) >= 0.1:
        if (dg == "bat_loi" and pi < pc) or (dg == "loi" and pi > pc):
            loi.append(f"{ten}: danh_gia «{dg}» ngược với tỷ lệ {pi}% (can thiệp) so với {pc}% (so sánh)")
    return loi


# ───────────────────────── KIỂM CHÉO ĐỘC LẬP (agent ngữ cảnh mới chấm TỪNG câu) ─────────────────────────
KET_LUAN_KIEM_CHEO = ("dung", "suy_luan_hop_ly", "ngoai_bai_co_nguon", "sai", "khong_co_trong_bai")
# 26/10/2026: bỏ «chinh», «danh_gia» khỏi tập này — kiểm chéo PMID 37142273 thấy `ket_cuc.chinh` (kết cục chính, chuỗi) và
# `nguy_co_sai_lech.mien[].danh_gia` (phán quyết từng miền RoB/AMSTAR-2) chưa từng được giao cho người chấm; mức «có một phần»
# không hợp lệ ở mục 11 AMSTAR-2 (mục chỉ có Có/Không) lọt qua, chỉ bị bắt gián tiếp qua `tong_the`. Giá trị không phải chuỗi
# (cờ `chinh: true` của dòng bảng) vẫn tự bị bỏ; dòng bảng ket_qua/tac_hai/khuyen_cao vẫn gộp thành MỘT mục như cũ.
_KHOA_META = {"phien_ban_khuon", "khoa", "pmid", "doi", "loai_thiet_ke", "may_tinh", "trang",
              "lien_quan_dashboard", "can_xac_nhan", "nguoi_soan", "trich_nguyen_van"}

HUONG_DAN_KIEM_CHEO = """# Kiểm chéo ĐỘC LẬP bản tóm tắt `tom_tat.json` với toàn văn `van_ban.txt`

Bạn KHÔNG soạn bản tóm tắt này. Bác sĩ sẽ dựa vào nó để điều trị người bệnh. Với TỪNG mục trong `kiem_cheo_can_kiem.json`
(đường dẫn + nội dung), đọc đúng chỗ trong `van_ban.txt` (trang mở đầu bằng `=== TRANG n ===`) rồi chấm một kết luận:
- `dung` — mọi dữ kiện và mọi con số trong câu có trong bài và gắn ĐÚNG đối tượng (đúng nhóm, đúng kết cục, đúng thời điểm);
- `suy_luan_hop_ly` — câu là thẩm định/diễn giải (nguy cơ sai lệch, GRADE, khả năng áp dụng) có căn cứ trong bài, không thêm
  dữ kiện mới;
- `ngoai_bai_co_nguon` — câu mang nhãn `[ngoài bài: …]` ghi rõ nguồn;
- `sai` — có dữ kiện/số SAI so với bài, gắn nhầm nhóm/kết cục, đảo chiều, hoặc thiếu mệnh đề điều kiện làm đổi nghĩa;
- `khong_co_trong_bai` — khẳng định không tìm thấy trong bài (kể cả so sánh nhất/kém nhất, nhận định về Việt Nam không ghi nguồn).
Mỗi mục ghi `trang` (danh sách trang đã đối chiếu) và `ghi_chu` (bắt buộc khi sai/khong_co_trong_bai: sai ở đâu, bài viết gì).
Không sửa tom_tat.json. Ghi kết quả vào `kiem_cheo.json` đúng khuôn:
{"phien_ban": 1, "tom_tat_sha256": "<sao y từ kiem_cheo_can_kiem.json>", "nguoi_kiem": "<tên agent, ngữ cảnh mới>",
 "luc": "<ISO>", "muc": [{"duong_dan": "...", "ket_luan": "dung", "trang": [3], "ghi_chu": ""}]}
"""


def cac_muc_kiem_cheo(tt: dict) -> list[tuple[str, str]]:
    """Mọi khẳng định phải được chấm: mỗi chuỗi chữ là một mục; mỗi dòng bảng (kết quả/tác hại/khuyến cáo) là MỘT mục."""
    ra: list[tuple[str, str]] = []

    def di(x, duong: str) -> None:
        if isinstance(x, dict):
            for k, v in x.items():
                if k in _KHOA_META:
                    continue
                if k == "danh_gia" and isinstance(v, str) and v.strip():
                    # phán quyết một miền đứng riêng không có nghĩa ⇒ gắn tên miền để người chấm đối chiếu đúng tiêu chí
                    ra.append((f"{duong}.danh_gia" if duong else "danh_gia",
                               f"{x.get('ten') or duong} → đánh giá: {v}"))
                    continue
                if k in ("ket_qua", "tac_hai", "khuyen_cao") and isinstance(v, list):
                    for i, h in enumerate(v, 1):
                        if isinstance(h, dict):
                            ra.append((f"{k}[{i}]", " | ".join(f"{a}={b}" for a, b in h.items() if a not in ("chinh",))))
                    continue
                di(v, f"{duong}.{k}" if duong else k)
        elif isinstance(x, list):
            for i, v in enumerate(x, 1):
                di(v, f"{duong}[{i}]")
        elif isinstance(x, str) and x.strip():
            ra.append((duong, x))
    di(tt, "")
    return ra


def bam_tom_tat(goi: Path) -> str:
    return hashlib.sha256((goi / "tom_tat.json").read_bytes()).hexdigest()


def chuan_bi_kiem_cheo(goi: Path) -> Path:
    tt = json.loads((goi / "tom_tat.json").read_text(encoding="utf-8"))
    ra = goi / "kiem_cheo_can_kiem.json"
    ra.write_text(json.dumps({"tom_tat_sha256": bam_tom_tat(goi),
                              "muc": [{"duong_dan": d, "noi_dung": n} for d, n in cac_muc_kiem_cheo(tt)]},
                             ensure_ascii=False, indent=1), encoding="utf-8")
    (goi / "HUONG-DAN-KIEM-CHEO.md").write_text(HUONG_DAN_KIEM_CHEO, encoding="utf-8")
    return ra


def kiem_ket_qua_kiem_cheo(goi: Path, tt: dict) -> list[str]:
    """Cổng kiểm chéo: phải có `kiem_cheo.json` của ĐÚNG bản tom_tat.json hiện tại (băm), chấm ĐỦ mọi mục, không mục nào sai."""
    tep = goi / "kiem_cheo.json"
    if not tep.is_file():
        return ["chưa có kiem_cheo.json — chạy `kiem-cheo <mã>` rồi giao một agent ĐỘC LẬP (ngữ cảnh mới) chấm từng câu"]
    try:
        kc = json.loads(tep.read_text(encoding="utf-8"))
    except ValueError as e:
        return [f"kiem_cheo.json không phải JSON hợp lệ: {e}"]
    loi: list[str] = []
    if kc.get("tom_tat_sha256") != bam_tom_tat(goi):
        return ["kiem_cheo.json chấm một bản tom_tat.json KHÁC (đã sửa sau khi kiểm chéo) — kiểm chéo lại bản hiện tại"]
    da_cham = {}
    for m in kc.get("muc") or []:
        if isinstance(m, dict):
            da_cham[m.get("duong_dan")] = m
    for duong, _ in cac_muc_kiem_cheo(tt):
        m = da_cham.get(duong)
        if m is None:
            loi.append(f"kiểm chéo bỏ sót mục {duong}")
            continue
        kl = m.get("ket_luan")
        if kl not in KET_LUAN_KIEM_CHEO:
            loi.append(f"kiểm chéo {duong}: kết luận «{kl}» không hợp lệ")
        elif kl in ("sai", "khong_co_trong_bai"):
            loi.append(f"kiểm chéo {duong}: {kl.upper()} — {str(m.get('ghi_chu') or '')[:200]}")
    return loi


# ───────────────────────── CỔNG kiểm tom_tat.json ─────────────────────────

def _rong(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return not v.strip() or v.strip() in ("…", "...")
    if isinstance(v, (list, dict)):
        return not v or all(_rong(x) for x in (v.values() if isinstance(v, dict) else v))
    return False


def _chuan_cong_cu(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower())


def _tap_so_trang(trang: list[str], so_trang: int) -> set[str]:
    ra: set[str] = set()
    for i in range(max(1, so_trang - 1), min(len(trang), so_trang + 1) + 1):
        ra.update(cac_so(trang[i - 1]))
    return ra


def _cau_co_tren_trang(cau: str, trang: list[str], so_trang: int) -> bool:
    """Câu trích có trên trang khai (±1) — khớp chuỗi sau chuẩn hoá, hoặc đủ các từ theo đúng thứ tự trong một khoảng
    ≤ 2× độ dài câu (PDF hay chèn xuống dòng/ký hiệu giữa chữ)."""
    q = _chuan_chu(cau).strip(" .\"'«»…")
    if not q:
        return False
    vung = _chuan_chu(" ".join(trang[max(0, so_trang - 2):min(len(trang), so_trang + 1)]))
    if q in vung:
        return True
    # PDF ngắt dòng có gạch nối kiểu «air -⏎way» hay «random-⏎ization»: so phần CHỈ chữ và số của hai bên.
    gon_q = re.sub(r"[\W_]+", "", q)
    if len(gon_q) >= 20 and gon_q in re.sub(r"[\W_]+", "", vung):
        return True
    tu_q = re.findall(r"\w+", q)
    tu_v = re.findall(r"\w+", vung)
    if not tu_q:
        return False
    for bat_dau in (i for i, w in enumerate(tu_v) if w == tu_q[0]):
        j, k = bat_dau, 0
        while j < len(tu_v) and j - bat_dau <= 2 * len(tu_q) and k < len(tu_q):
            if tu_v[j] == tu_q[k]:
                k += 1
            j += 1
        if k == len(tu_q):
            return True
    return False


# Bác sĩ 19/10/2026: «Loại bỏ những câu chữ cần kiểm chứng và cần xác nhận; nếu cần thì để cuối văn bản» + «những câu đề
# xuất của máy cũng bỏ đi». Thân văn bản không được chứa các cụm dưới; đưa vào `can_xac_nhan` (in ở mục Lưu ý cuối).
CUM_CAM_THAN_BAI = re.compile(r"\[\s*CẦN|cần\s+(?:bác\s+sĩ\s+)?(?:kiểm\s+chứng|xác\s+nhận|tra\s+thêm)|đề\s+xuất\s+(?:của\s+)?máy"
                              r"|máy\s+đề\s+xuất|chưa\s+qua\s+bác\s+sĩ|đề\s+xuất\s+mức", re.I)
_KHONG_KIEM_CUM = {"can_xac_nhan", "nguoi_soan"}


def cac_chuoi_than_bai(x, duong: str = "") -> list[tuple[str, str]]:
    ra: list[tuple[str, str]] = []
    if isinstance(x, dict):
        for k, v in x.items():
            if k not in _KHONG_KIEM_CUM:
                ra += cac_chuoi_than_bai(v, f"{duong}.{k}" if duong else k)
    elif isinstance(x, list):
        for i, v in enumerate(x, 1):
            ra += cac_chuoi_than_bai(v, f"{duong}[{i}]")
    elif isinstance(x, str):
        ra.append((duong, x))
    return ra


_KHONG_KIEM_SO = {"phien_ban_khuon", "khoa", "pmid", "doi", "nguoi_soan", "trich_nguyen_van", "ket_qua", "tac_hai", "trang",
                  "dang_ky", "cong_cu"}  # cong_cu: tên «RoB 2», «AMSTAR-2» không phải số liệu


def cac_chuoi_kiem_so(x, duong: str = "") -> list[tuple[str, str]]:
    """Mọi chuỗi chữ của tom_tat.json phải qua phép kiểm số — trừ định danh, bảng đã kiểm theo trang, số trang."""
    ra: list[tuple[str, str]] = []
    if isinstance(x, dict):
        for k, v in x.items():
            if k not in _KHONG_KIEM_SO:
                ra += cac_chuoi_kiem_so(v, f"{duong}.{k}" if duong else k)
    elif isinstance(x, list):
        for i, v in enumerate(x, 1):
            ra += cac_chuoi_kiem_so(v, f"{duong}[{i}]")
    elif isinstance(x, str):
        ra.append((duong, x))
    return ra


def kiem_tom_tat(tt: dict, trang: list[str], ngu_canh: dict) -> list[str]:
    """Danh sách lỗi (rỗng = đạt). Mỗi lỗi là một câu tiếng Việt chỉ đúng chỗ phải sửa."""
    loi: list[str] = []
    if not isinstance(tt, dict):
        return ["tom_tat.json không phải đối tượng JSON"]
    if tt.get("phien_ban_khuon") != PHIEN_BAN_KHUON:
        loi.append(f"phien_ban_khuon phải là {PHIEN_BAN_KHUON}")
    bg = ngu_canh.get("ban_ghi") or {}
    if tt.get("khoa") != bg.get("khoa"):
        loi.append(f"khoa «{tt.get('khoa')}» khác bản ghi «{bg.get('khoa')}»")
    if bg.get("pmid") and str(tt.get("pmid") or "") != bg["pmid"]:
        loi.append(f"pmid «{tt.get('pmid')}» khác PMID của bản ghi EndNote «{bg['pmid']}»")
    tk = tt.get("loai_thiet_ke")
    if tk not in LOAI_THIET_KE:
        loi.append(f"loai_thiet_ke phải thuộc {LOAI_THIET_KE}")
        return loi
    for k in _BAT_BUOC_CHUNG + _BAT_BUOC_THEO_THIET_KE[tk]:
        if _rong(tt.get(k)):
            loi.append(f"thiếu hoặc rỗng: {k}")
    pico = tt.get("cau_hoi_lam_sang") or {}
    for k in ("loai", "P", "O"):
        if _rong(pico.get(k)):
            loi.append(f"cau_hoi_lam_sang.{k} rỗng")
    kl = tt.get("ket_luan_thuc_hanh") or []
    if not isinstance(kl, list) or not 2 <= len([x for x in kl if not _rong(x)]) <= 8:
        loi.append("ket_luan_thuc_hanh cần 2–8 gạch có nội dung")
    goi_y = tt.get("goi_y_ap_dung") or {}
    if _rong(goi_y.get("nen_lam")) and _rong(goi_y.get("khong_nen")):
        loi.append("goi_y_ap_dung cần ít nhất «nen_lam» hoặc «khong_nen»")
    # Công cụ nguy cơ sai lệch đúng thiết kế
    nc_sl = tt.get("nguy_co_sai_lech") or {}
    cc = _chuan_cong_cu(nc_sl.get("cong_cu", ""))
    hop_le = CONG_CU_HOP_LE[tk]
    if hop_le and not any(cc.startswith(h) for h in hop_le):
        loi.append(f"nguy_co_sai_lech.cong_cu «{nc_sl.get('cong_cu')}» không đúng thiết kế {tk} (hợp lệ: {', '.join(hop_le)})")
    if tk != "rct" and cc.startswith("rob2"):
        loi.append("RoB 2 chỉ dùng cho RCT")
    if tk not in ("nhan_thuoc",) and _rong(nc_sl.get("mien")):
        loi.append("nguy_co_sai_lech.mien rỗng — đánh giá từng miền của công cụ")
    for i, mi in enumerate(nc_sl.get("mien") or [], 1):
        if not isinstance(mi, dict) or _rong(mi.get("ten")) or _rong(mi.get("danh_gia")) or _rong(mi.get("ly_do")):
            loi.append(f"nguy_co_sai_lech.mien[{i}] cần ten · danh_gia · ly_do")
    # GRADE đề xuất
    gr = tt.get("grade_de_xuat") or []
    for i, g in enumerate(gr if isinstance(gr, list) else [], 1):
        if not isinstance(g, dict) or g.get("muc") not in MUC_GRADE or _rong(g.get("ly_do")) or _rong(g.get("ket_cuc")):
            loi.append(f"grade_de_xuat[{i}]: cần ket_cuc · muc ∈ {MUC_GRADE} · ly_do")
    # SỐ LIỆU: mọi số có nguyên văn trên trang khai (±1)
    for khoi in ("ket_qua", "tac_hai"):
        hang = tt.get(khoi) or []
        if not isinstance(hang, list):
            loi.append(f"{khoi} phải là danh sách")
            continue
        for i, h in enumerate(hang, 1):
            if not isinstance(h, dict) or _rong(h.get("ket_cuc")):
                loi.append(f"{khoi}[{i}] thiếu ket_cuc")
                continue
            tr = h.get("trang")
            if not isinstance(tr, int) or not 1 <= tr <= len(trang):
                loi.append(f"{khoi}[{i}] «{h.get('ket_cuc')}»: trang {tr!r} không hợp lệ (PDF có {len(trang)} trang)")
                continue
            tap = _tap_so_trang(trang, tr)
            may_tinh = set(h.get("may_tinh") or [])
            for truong in TRUONG_SO:
                if truong in may_tinh:
                    continue
                for so in cac_so(str(h.get(truong) or "")):
                    if not _so_khop(so, tap):
                        loi.append(f"{khoi}[{i}] «{h.get('ket_cuc')}».{truong}: số {so} KHÔNG có trên trang {tr}±1 "
                                   "— chép sai, sai trang, hay do máy tính (khai may_tinh)?")
    if tk in ("guideline", "nhan_thuoc"):
        for i, kc in enumerate(tt.get("khuyen_cao") or [], 1):
            if not isinstance(kc, dict) or _rong(kc.get("noi_dung")):
                loi.append(f"khuyen_cao[{i}] thiếu noi_dung")
                continue
            tr = kc.get("trang")
            if not isinstance(tr, int) or not 1 <= tr <= len(trang):
                loi.append(f"khuyen_cao[{i}]: trang {tr!r} không hợp lệ")
    # ĐIỂM NHẤN từng mục
    dn = tt.get("diem_nhan") if isinstance(tt.get("diem_nhan"), dict) else {}
    for k in KHOA_DIEM_NHAN:
        if k == "tac_hai" and tk in ("guideline", "nhan_thuoc"):
            continue
        if _rong(dn.get(k)):
            loi.append(f"diem_nhan.{k} rỗng — mỗi mục của Word mở đầu bằng một ô ĐIỂM NHẤN")
        elif len(re.findall(r"\S+", str(dn.get(k)))) > TOI_DA_TU_DIEM_NHAN:
            loi.append(f"diem_nhan.{k} dài quá {TOI_DA_TU_DIEM_NHAN} từ — điểm nhấn phải ngắn")
    for i, h in enumerate(tt.get("tac_hai") or [], 1):
        if isinstance(h, dict) and h.get("danh_gia") not in (None, "") + DANH_GIA_TAC_HAI:
            loi.append(f"tac_hai[{i}].danh_gia phải thuộc {DANH_GIA_TAC_HAI}")
    # SỐ ở MỌI trường chữ (kết luận, điểm nhấn, gợi ý, phương pháp, nguy cơ sai lệch, GRADE, điểm mạnh/hạn chế…): phải có
    # nguyên văn trong bài hoặc trong bảng đã kiểm theo trang. Đo 19/10 (PARADIGM-HF): tổng tự cộng «8399», «2031» lọt vào
    # điểm nhấn/điểm mạnh khi cổng chỉ kiểm phần tóm lược. Chuỗi mang nhãn «[ngoài bài: …]» được miễn (Word in xám nghiêng).
    tap_bai: set[str] = set()
    for t_ in trang:
        tap_bai.update(cac_so(t_))
    tap_bang: set[str] = set()
    for h in (tt.get("ket_qua") or []) + (tt.get("tac_hai") or []):
        if isinstance(h, dict):
            for truong in TRUONG_SO:
                tap_bang.update(cac_so(str(h.get(truong) or "")))
    for ten, x in cac_chuoi_kiem_so(tt):
        x = str(x or "")
        if "[ngoài bài" in x:
            continue
        sai = [so for so in cac_so(x) if not (_so_khop(so, tap_bai) or _so_khop(so, tap_bang))]
        if sai:
            loi.append(f"{ten}: số {', '.join(sorted(set(sai)))} không có trong bài — chép đúng số của bài, hoặc ghi "
                       "«[ngoài bài: <nguồn>]» nếu lấy từ nguồn khác")
    for khoi in ("ket_qua", "tac_hai"):
        for i, h in enumerate(tt.get(khoi) or [], 1):
            if isinstance(h, dict):
                loi += kiem_nhat_quan_dong(h, f"{khoi}[{i}] «{h.get('ket_cuc')}»")
    for ten, x in cac_chuoi_than_bai(tt):
        g = SO_GAN_DUNG.search(x)
        if g:
            loi.append(f"{ten}: số gần đúng «{g.group(0)}» — ghi đúng số của bài (từng nhóm nếu bài báo cáo theo nhóm)")
        m = CUM_CAM_THAN_BAI.search(x)
        if m:
            loi.append(f"{ten}: thân văn bản chứa «{m.group(0)}» — chuyển điểm chưa chắc vào can_xac_nhan (in ở cuối văn bản)")
    if not isinstance(tt.get("can_xac_nhan", []), list):
        loi.append("can_xac_nhan phải là danh sách")
    # Trích nguyên văn (tuỳ chọn, không in ra Word — vẫn kiểm nếu có)
    for i, tc in enumerate(tt.get("trich_nguyen_van") or [], 1):
        cau = (tc or {}).get("cau", "") if isinstance(tc, dict) else ""
        tr = (tc or {}).get("trang") if isinstance(tc, dict) else None
        if len(re.findall(r"\S+", cau)) > TOI_DA_TU_TRICH:
            loi.append(f"trich_nguyen_van[{i}] dài {len(cau.split())} từ (tối đa {TOI_DA_TU_TRICH})")
        if not isinstance(tr, int) or not 1 <= tr <= len(trang) or not _cau_co_tren_trang(cau, trang, tr):
            loi.append(f"trich_nguyen_van[{i}] «{cau[:60]}» không tìm thấy nguyên văn trên trang {tr}±1")
    return loi


# ───────────────────────── dựng Word (python-docx, Times New Roman, có ĐIỂM NHẤN màu) ─────────────────────────
# Bác sĩ 19/10/2026: «Phần 12, 13, 14 không cần, mỗi phần trình bày phải có điểm nhấn với những in đậm, màu sắc».
# Màu mang NGHĨA, dùng nhất quán: xanh lá = lợi ích / nguy cơ sai lệch thấp / nên làm · đỏ = cảnh báo / nguy cơ cao /
# không nên · cam = thận trọng / cần xác nhận · xanh dương = theo dõi / thông tin · xám = ngoài bài, chú thích.
# Không dùng ký hiệu ngoài bảng mã Times New Roman (☐ ✔ ★… in ra ô vuông rỗng) — nhấn bằng NHÃN CHỮ IN HOA + màu.

MAU_CHU = {"navy": "1F3864", "duong": "2E74B5", "la": "1E7B34", "do": "C00000", "cam": "B45F06", "xam": "595959",
           "trang": "FFFFFF"}
MAU_NEN = {"navy": "1F3864", "duong": "DEEAF6", "la": "E2F0D9", "do": "FBE4E4", "cam": "FFF2CC", "xam": "F2F2F2"}
_DANH_DAU = re.compile(r"\*\*(.+?)\*\*|\[!(.+?)!\]|\[\+(.+?)\+\]|(\[CẦN[^\]]*\])|(\[ngoài bài[^\]]*\])")
_SAU_PBDR = ("shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
             "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
             "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId",
             "cnfStyle", "rPr", "sectPr", "pPrChange")
_SAU_SHD_O = ("noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark")


def _o(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip() or "—"


def _ox():
    from docx.oxml import OxmlElement  # noqa: PLC0415
    from docx.oxml.ns import qn  # noqa: PLC0415
    return OxmlElement, qn


def _font(run, *, dam: bool = False, nghieng: bool = False, mau: str | None = None, co: float | None = None) -> None:
    from docx.shared import Pt, RGBColor  # noqa: PLC0415
    OxmlElement, qn = _ox()
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), "Times New Roman")
    run.bold = dam or None
    run.italic = nghieng or None
    if mau:
        run.font.color.rgb = RGBColor.from_string(mau)
    if co:
        run.font.size = Pt(co)


def _them(par, text: str, **kieu):
    r = par.add_run(text)
    _font(r, **kieu)
    return r


def _viet(par, text, *, co: float | None = None, mau: str | None = None, dam: bool = False) -> None:
    """Chữ có đánh dấu nhấn: **đậm** · [!cảnh báo — đỏ đậm!] · [+lợi ích — xanh lá đậm+] · [CẦN …] cam đậm ·
    [ngoài bài …] xám nghiêng (số liệu không lấy từ bài này)."""
    text = _o(text)
    vi_tri = 0
    for m in _DANH_DAU.finditer(text):
        if m.start() > vi_tri:
            _them(par, text[vi_tri:m.start()], co=co, mau=mau, dam=dam)
        if m.group(1):
            _them(par, m.group(1), co=co, mau=mau or MAU_CHU["navy"], dam=True)
        elif m.group(2):
            _them(par, m.group(2), co=co, mau=MAU_CHU["do"], dam=True)
        elif m.group(3):
            _them(par, m.group(3), co=co, mau=MAU_CHU["la"], dam=True)
        elif m.group(4):
            _them(par, m.group(4), co=co, mau=MAU_CHU["cam"], dam=True)
        else:
            _them(par, m.group(5), co=co, mau=MAU_CHU["xam"], nghieng=True)
        vi_tri = m.end()
    if vi_tri < len(text):
        _them(par, text[vi_tri:], co=co, mau=mau, dam=dam)


def _chen_dung_thu_tu(cha, con, sau: tuple[str, ...]) -> None:
    """Chèn phần tử OOXML TRƯỚC phần tử đầu tiên thuộc nhóm phải đứng sau nó (CLAUDE.md: pBdr trước spacing, shd sau
    tcBorders) — append bừa làm Word báo tệp hỏng."""
    _, qn = _ox()
    for el in cha:
        if el.tag in {qn(f"w:{t}") for t in sau}:
            el.addprevious(con)
            return
    cha.append(con)


def _to_nen(o, mau_nen: str) -> None:
    OxmlElement, qn = _ox()
    tcpr = o._tc.get_or_add_tcPr()
    cu = tcpr.find(qn("w:shd"))
    if cu is not None:
        tcpr.remove(cu)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), mau_nen)
    _chen_dung_thu_tu(tcpr, shd, _SAU_SHD_O)


def _doan(noi, *, sau: float = 3, truoc: float = 0, can: str | None = None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: PLC0415
    from docx.shared import Pt  # noqa: PLC0415
    p = noi.add_paragraph()
    p.paragraph_format.space_after = Pt(sau)
    p.paragraph_format.space_before = Pt(truoc)
    if can == "giua":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _tieu_de_muc(doc, so: int, ten: str) -> None:
    OxmlElement, qn = _ox()
    from docx.shared import Pt  # noqa: PLC0415
    p = doc.add_heading(level=2)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    _them(p, f"{so}. {ten}", dam=True, mau=MAU_CHU["navy"], co=14)
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    duoi = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", MAU_CHU["duong"])):
        duoi.set(qn(k), v)
    bdr.append(duoi)
    _chen_dung_thu_tu(ppr, bdr, _SAU_PBDR)


def _tieu_de_nho(doc, ten: str, mau: str = "duong") -> None:
    p = _doan(doc, truoc=6, sau=2)
    _them(p, ten, dam=True, mau=MAU_CHU[mau], co=12.5)


def _hop(doc, nhan: str, noi_dung, *, mau: str = "duong", danh_sach: bool = False) -> None:
    """Ô tô nền một ô: NHÃN in hoa đậm màu + nội dung (chuỗi hoặc danh sách gạch đánh số)."""
    from docx.shared import Pt  # noqa: PLC0415
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.autofit = False
    o = t.cell(0, 0)
    _dat_rong(t, [17.0])
    _to_nen(o, MAU_NEN[mau])
    p = o.paragraphs[0]
    _them(p, nhan, dam=True, mau=MAU_CHU["navy" if mau == "duong" else mau])
    if danh_sach:
        for i, x in enumerate([x for x in noi_dung or [] if not _rong(x)], 1):
            q = o.add_paragraph()
            q.paragraph_format.space_after = Pt(3)
            _them(q, f"{i}. ", dam=True, mau=MAU_CHU["navy" if mau == "duong" else mau])
            _viet(q, x)
    else:
        _them(p, "  ")
        _viet(p, noi_dung)
    _doan(doc, sau=2)


def _bang(doc, tieu_de: list[str], hang: list[list], *, co: float = 10.5, rong: list[float] | None = None) -> None:
    """Bảng có dòng tiêu đề nền xanh đậm chữ trắng (lặp lại khi sang trang). Ô = chuỗi, hoặc (chuỗi, {dam, mau, nen}).
    `rong` = bề rộng từng cột (cm); vắng ⇒ chia đều 17 cm (trọn bề ngang A4 lề 2 cm)."""
    OxmlElement, qn = _ox()
    t = doc.add_table(rows=1, cols=len(tieu_de))
    t.style = "Table Grid"
    t.autofit = False
    rong = rong or [17.0 / len(tieu_de)] * len(tieu_de)
    dau = t.rows[0]
    trpr = dau._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for i, h in enumerate(tieu_de):
        _to_nen(dau.cells[i], MAU_NEN["navy"])
        _them(dau.cells[i].paragraphs[0], h, dam=True, mau=MAU_CHU["trang"], co=co)
    for h in hang:
        o = t.add_row().cells
        for i, gt in enumerate(h):
            kieu = {}
            if isinstance(gt, tuple):
                gt, kieu = gt
            if kieu.get("nen"):
                _to_nen(o[i], MAU_NEN[kieu["nen"]])
            _viet(o[i].paragraphs[0], gt, co=co, mau=MAU_CHU.get(kieu.get("mau", ""), None), dam=kieu.get("dam", False))
    _dat_rong(t, rong)
    _doan(doc, sau=2)


def _dat_rong(t, rong: list[float]) -> None:
    """Bề rộng bảng ở CẢ ba nơi Word/trình xem đọc: lưới cột (tblGrid), từng ô (tcW) và tổng bảng (tblW, kiểu dxa) —
    chỉ đặt tcW thì Xem nhanh của macOS (và vài bản Word) vẫn vẽ theo lưới cột mặc định, bảng hẹp lệch trái."""
    from docx.shared import Cm  # noqa: PLC0415
    OxmlElement, qn = _ox()
    for i, cot in enumerate(t.columns):
        cot.width = Cm(rong[i])
    for dong in t.rows:
        for i, o_ in enumerate(dong.cells):
            o_.width = Cm(rong[i])
    tblpr = t._tbl.tblPr
    w = tblpr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW")
        tblpr.append(w)
    w.set(qn("w:w"), str(int(sum(rong) / 2.54 * 1440)))
    w.set(qn("w:type"), "dxa")


def _danh_sach(doc, xs, *, mau_dau: str = "navy") -> None:
    for x in [x for x in xs or [] if not _rong(x)]:
        p = _doan(doc, sau=2)
        p.paragraph_format.left_indent = 270000  # ≈ 0.75 cm
        p.paragraph_format.first_line_indent = -150000
        _them(p, "•  ", dam=True, mau=MAU_CHU[mau_dau])
        _viet(p, x)


def _nhan_dong(doc, nhan: str, noi_dung, mau: str) -> None:
    p = _doan(doc, sau=3)
    _them(p, f"{nhan}: ", dam=True, mau=MAU_CHU[mau])
    _viet(p, noi_dung)


def _mau_danh_gia(s: str, *, agree: bool = False) -> str | None:
    """Màu ô đánh giá nguy cơ sai lệch. AGREE II đo CHẤT LƯỢNG (cao = tốt) ⇒ đảo chiều so với RoB/ROBINS/QUADAS."""
    s = _chuan_chu(s)
    cao = any(k in s for k in ("cao", "high", "nghiêm trọng", "serious", "critical"))
    lo = any(k in s for k in ("lo ngại", "concern", "trung bình", "moderate", "không rõ", "unclear", "vừa"))
    thap = any(k in s for k in ("thấp", "low"))
    if lo:
        return "cam"
    if cao and not thap:
        return "la" if agree else "do"
    if thap and not cao:
        return "do" if agree else "la"
    return None


RONG_BANG_SO = [4.1, 2.3, 2.3, 1.9, 1.9, 1.3, 1.4, 1.1, 0.7]  # cm — bảng kết quả/tác hại 9 cột, tổng 17 cm
MAU_GRADE = {"cao": "la", "trung bình": "duong", "thấp": "cam", "rất thấp": "do", "không áp dụng": "xam"}
MAU_TAC_HAI = {"bat_loi": "do", "loi": "la", "tuong_duong": "xam"}


def _gia_tri(h: dict, truong: str) -> str:
    v = _o(h.get(truong))
    return f"{v} †" if v != "—" and truong in set(h.get("may_tinh") or []) else v


def dung_word(tt: dict, ngu_canh: dict):
    """Dựng Document từ tom_tat.json ĐÃ QUA CỔNG. 11 mục (guideline/nhãn thuốc: 10 — không có mục tác hại riêng)."""
    from docx import Document  # noqa: PLC0415
    from docx.shared import Cm, Pt  # noqa: PLC0415
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for k in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, k, Cm(2))
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    OxmlElement, qn = _ox()
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), "Times New Roman")
    zoom = doc.settings.element.find(qn("w:zoom"))  # mẫu python-docx thiếu w:percent bắt buộc — trình xác thực OOXML báo lỗi
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    bg = ngu_canh.get("ban_ghi") or {}
    pdf = ngu_canh.get("pdf") or {}
    rb = ngu_canh.get("rut_bai") or {}
    tk = tt["loai_thiet_ke"]
    dn = tt.get("diem_nhan") or {}
    la_guideline = tk in ("guideline", "nhan_thuoc")

    # Đầu trang
    _them(_doan(doc, sau=4, can="giua"), "TÓM TẮT CHỨNG CỨ CHO THỰC HÀNH LÂM SÀNG", dam=True, mau=MAU_CHU["navy"], co=16)
    _them(_doan(doc, sau=3, can="giua"), _o(bg.get("tieu_de")), dam=True, co=13)
    _them(_doan(doc, sau=6, can="giua"),
          f"{_o(bg.get('tac_gia'))} · {_o(bg.get('tap_chi'))} · {_o(bg.get('nam'))}"
          + (f" · PMID {bg['pmid']}" if bg.get("pmid") else "") + (f" · DOI {bg['doi']}" if bg.get("doi") else ""),
          nghieng=True, mau=MAU_CHU["xam"], co=11)
    p = _doan(doc, sau=8, can="giua")
    _them(p, "THIẾT KẾ: ", dam=True, mau=MAU_CHU["navy"], co=11)
    _them(p, TEN_THIET_KE[tk], co=11)
    _them(p, "     RÚT BÀI: ", dam=True, mau=MAU_CHU["navy"], co=11)
    tt_rb = rb.get("trang_thai")
    _them(p, {"khong_thay_rut": "không thấy thông báo rút", "khong_thay_rut_rw": "không thấy (chỉ Retraction Watch)",
              "da_rut": "ĐÃ RÚT", "chua_kiem": "CHƯA KIỂM ĐƯỢC"}.get(tt_rb, "—"),
          dam=True, mau=MAU_CHU["la" if tt_rb == "khong_thay_rut" else ("cam" if tt_rb == "khong_thay_rut_rw" else "do")], co=11)

    so = 1
    _tieu_de_muc(doc, so, "Kết luận cho thực hành")
    _hop(doc, "ĐIỂM NHẤN CHO THỰC HÀNH", tt.get("ket_luan_thuc_hanh"), mau="la", danh_sach=True)

    so += 1
    _tieu_de_muc(doc, so, "Thông tin trích dẫn")
    _hop(doc, "ĐIỂM NHẤN:", dn.get("thong_tin"))
    bc = tt.get("boi_canh") or {}
    _bang(doc, ["Mục", "Nội dung"], [
        [("Thiết kế", {"dam": True, "nen": "xam"}), f"**{TEN_THIET_KE[tk]}** — {_o(tt.get('thiet_ke_mo_ta'))}"],
        [("Lĩnh vực", {"dam": True, "nen": "xam"}), tt.get("linh_vuc")],
        [("Địa điểm", {"dam": True, "nen": "xam"}), bc.get("dia_diem")],
        [("Thời gian", {"dam": True, "nen": "xam"}), bc.get("thoi_gian")],
        [("Đăng ký", {"dam": True, "nen": "xam"}), bc.get("dang_ky")],
        [("Tài trợ", {"dam": True, "nen": "xam"}), bc.get("tai_tro")],
        [("Xung đột lợi ích", {"dam": True, "nen": "xam"}), bc.get("xung_dot_loi_ich")],
        [("Rút bài", {"dam": True, "nen": "xam"}),
         (f"{NHAN_RUT_BAI.get(tt_rb, '—')} — tra {_o(rb.get('luc'))} ({_o(rb.get('nguon'))})",
          {"mau": "la" if tt_rb == "khong_thay_rut" else "do", "dam": tt_rb != "khong_thay_rut"})],
    ], co=11, rong=[4.0, 13.0])

    so += 1
    pico = tt.get("cau_hoi_lam_sang") or {}
    _tieu_de_muc(doc, so, "Câu hỏi lâm sàng (PICO)")
    _hop(doc, "ĐIỂM NHẤN:", dn.get("pico"))
    _bang(doc, ["Thành phần", "Nội dung"], [
        [("Loại câu hỏi", {"dam": True, "nen": "xam"}), pico.get("loai")],
        [("P — Người bệnh", {"dam": True, "mau": "navy", "nen": "duong"}), pico.get("P")],
        [("I — Can thiệp / phơi nhiễm", {"dam": True, "mau": "navy", "nen": "duong"}), pico.get("I")],
        [("C — So sánh", {"dam": True, "mau": "navy", "nen": "duong"}), pico.get("C")],
        [("O — Kết cục", {"dam": True, "mau": "navy", "nen": "duong"}), pico.get("O")]], co=11, rong=[4.5, 12.5])

    so += 1
    if la_guideline:
        _tieu_de_muc(doc, so, "Phạm vi và phương pháp xây dựng")
        _hop(doc, "ĐIỂM NHẤN:", dn.get("phuong_phap"))
        _viet(_doan(doc), tt.get("phan_tich"))
        dt_ = tt.get("doi_tuong") or {}
        if not _rong(dt_):
            _nhan_dong(doc, "Đối tượng áp dụng", dt_.get("dac_diem_nen") or dt_.get("co_mau"), "navy")
        so += 1
        _tieu_de_muc(doc, so, "Các khuyến cáo (phân hạng NGUYÊN BẢN của nguồn)")
        _hop(doc, "ĐIỂM NHẤN:", dn.get("ket_qua"))
        _bang(doc, ["#", "Khuyến cáo", "Lực khuyến cáo", "Mức chứng cứ", "Trang"],
              [[str(i), (kc.get("noi_dung"), {"dam": bool(kc.get("lien_quan_dashboard"))}),
                (kc.get("luc_khuyen_cao"), {"dam": True, "mau": "navy"}), (kc.get("muc_chung_cu"), {"dam": True}),
                kc.get("trang")] for i, kc in enumerate(tt.get("khuyen_cao") or [], 1)], co=10.5,
              rong=[0.8, 9.6, 2.9, 2.7, 1.0])
    else:
        dt_ = tt.get("doi_tuong") or {}
        ct = tt.get("can_thiep_so_sanh") or {}
        kc_ = tt.get("ket_cuc") or {}
        _tieu_de_muc(doc, so, "Phương pháp")
        _hop(doc, "ĐIỂM NHẤN:", dn.get("phuong_phap"))
        _tieu_de_nho(doc, "Đối tượng")
        _nhan_dong(doc, "Cỡ mẫu", dt_.get("co_mau"), "navy")
        _them(_doan(doc, sau=1), "Tiêu chuẩn chọn", dam=True, mau=MAU_CHU["la"])
        _danh_sach(doc, dt_.get("tieu_chuan_chon"), mau_dau="la")
        _them(_doan(doc, sau=1), "Tiêu chuẩn loại", dam=True, mau=MAU_CHU["do"])
        _danh_sach(doc, dt_.get("tieu_chuan_loai"), mau_dau="do")
        _nhan_dong(doc, "Đặc điểm nền", dt_.get("dac_diem_nen"), "navy")
        if not _rong(ct):
            _tieu_de_nho(doc, "Can thiệp và so sánh")
            _nhan_dong(doc, "Can thiệp", ct.get("can_thiep"), "la")
            _nhan_dong(doc, "So sánh", ct.get("so_sanh"), "duong")
            _nhan_dong(doc, "Điều trị đi kèm", ct.get("dong_tri_lieu"), "xam")
        _tieu_de_nho(doc, "Kết cục và theo dõi")
        _nhan_dong(doc, "Kết cục chính", kc_.get("chinh"), "navy")
        _them(_doan(doc, sau=1), "Kết cục phụ", dam=True, mau=MAU_CHU["duong"])
        _danh_sach(doc, kc_.get("phu"), mau_dau="duong")
        _nhan_dong(doc, "Thời gian theo dõi", kc_.get("theo_doi"), "navy")
        _tieu_de_nho(doc, "Phân tích thống kê")
        _viet(_doan(doc), tt.get("phan_tich"))

        cot = ["Kết cục", "Can thiệp", "So sánh", "Hiệu ứng", "KTC 95%", "p", "Tuyệt đối", "NNT/NNH", "Tr."]

        def _dong(h: dict, *, nen: str | None = None) -> list:
            ten = _o(h.get("ket_cuc")) + (" (CHÍNH)" if h.get("chinh") else "")
            nen = nen or ("la" if h.get("chinh") else None)
            o = [(ten, {"dam": bool(h.get("chinh")), "nen": nen})]
            for t_ in TRUONG_SO:
                o.append((_gia_tri(h, t_), {"dam": t_ in ("hieu_ung", "ktc95"), "nen": nen,
                                            "mau": "navy" if t_ in ("hieu_ung", "ktc95") else ""}))
            return o + [(_o(h.get("trang")), {"nen": nen})]
        so += 1
        _tieu_de_muc(doc, so, "Kết quả")
        _hop(doc, "ĐIỂM NHẤN:", dn.get("ket_qua"), mau="la")
        _bang(doc, cot, [_dong(h) for h in tt.get("ket_qua") or []], co=10, rong=RONG_BANG_SO)
        if any(h.get("may_tinh") for h in (tt.get("ket_qua") or []) + (tt.get("tac_hai") or [])):
            _them(_doan(doc), "† Số do MÁY tính từ số liệu của bài (cách tính ở mục Phân tích thống kê) — không phải số bài "
                  "báo công bố. Dòng nền xanh lá = kết cục chính.", nghieng=True, mau=MAU_CHU["xam"], co=10)
        if not _rong(tt.get("phan_nhom")):
            _them(_doan(doc, sau=1), "Phân nhóm", dam=True, mau=MAU_CHU["duong"])
            _danh_sach(doc, tt.get("phan_nhom"), mau_dau="duong")
        so += 1
        _tieu_de_muc(doc, so, "Tác hại và an toàn")
        _hop(doc, "ĐIỂM NHẤN:", dn.get("tac_hai"), mau="do")
        if tt.get("tac_hai"):
            _bang(doc, cot, [_dong(h, nen=MAU_TAC_HAI.get(h.get("danh_gia") or "")) for h in tt["tac_hai"]], co=10,
                  rong=RONG_BANG_SO)
            _them(_doan(doc), "Nền đỏ = nhóm can thiệp nhiều tác hại hơn · nền xanh lá = ít hơn · nền xám = không khác biệt.",
                  nghieng=True, mau=MAU_CHU["xam"], co=10)
        if not _rong(tt.get("tac_hai_ghi_chu")):
            _viet(_doan(doc), tt.get("tac_hai_ghi_chu"))

    so += 1
    nsl = tt.get("nguy_co_sai_lech") or {}
    agree = _chuan_cong_cu(nsl.get("cong_cu", "")).startswith("agree")
    _tieu_de_muc(doc, so, f"Thẩm định nguy cơ sai lệch — {_o(nsl.get('cong_cu'))}")
    _hop(doc, "ĐIỂM NHẤN:", dn.get("nguy_co_sai_lech"))
    _bang(doc, ["Miền", "Đánh giá", "Lý do", "Tr."],
          [[(m.get("ten"), {"dam": True}),
            (m.get("danh_gia"), {"dam": True, "nen": _mau_danh_gia(str(m.get("danh_gia")), agree=agree)}),
            m.get("ly_do"), m.get("trang")] for m in nsl.get("mien") or []], co=10.5, rong=[3.6, 2.6, 9.8, 1.0])
    p = _doan(doc)
    _them(p, "TỔNG THỂ: ", dam=True, mau=MAU_CHU[_mau_danh_gia(str(nsl.get("tong_the")), agree=agree) or "navy"])
    _viet(p, nsl.get("tong_the"))

    so += 1
    so_grade = so
    _tieu_de_muc(doc, so, "Mức chắc chắn của chứng cứ (GRADE)")
    _hop(doc, "ĐIỂM NHẤN:", dn.get("grade"))
    _bang(doc, ["Kết cục", "Mức chắc chắn", "Lý do hạ / nâng"],
          [[(g.get("ket_cuc"), {"dam": True}),
            (str(g.get("muc", "")).upper(), {"dam": True, "nen": MAU_GRADE.get(g.get("muc"))}),
            g.get("ly_do")] for g in tt.get("grade_de_xuat") or []], co=10.5, rong=[4.2, 2.8, 10.0])

    so += 1
    ka = tt.get("kha_nang_ap_dung") or {}
    _tieu_de_muc(doc, so, "Khả năng áp dụng cho người bệnh ngoại trú Việt Nam")
    _hop(doc, "ĐIỂM NHẤN:", dn.get("ap_dung_vn"))
    _nhan_dong(doc, "TƯƠNG ĐỒNG", ka.get("tuong_dong"), "la")
    _nhan_dong(doc, "KHÁC BIỆT", ka.get("khac_biet"), "cam")
    _nhan_dong(doc, "TRIỂN KHAI TẠI VIỆT NAM", ka.get("trien_khai_viet_nam"), "duong")

    so += 1
    gy = tt.get("goi_y_ap_dung") or {}
    _tieu_de_muc(doc, so, "Áp dụng vào thực hành")
    _hop(doc, "ĐIỂM NHẤN:", dn.get("goi_y"), mau="la")
    for nhan, khoa, mau in (("NÊN LÀM", "nen_lam", "la"), ("KHÔNG NÊN", "khong_nen", "do"),
                            ("THẬN TRỌNG / ĐIỀU KIỆN", "dieu_kien", "cam"), ("THEO DÕI", "theo_doi", "duong")):
        if not _rong(gy.get(khoa)):
            _them(_doan(doc, truoc=4, sau=1), nhan, dam=True, mau=MAU_CHU[mau])
            _danh_sach(doc, gy.get(khoa), mau_dau=mau)

    so += 1
    _tieu_de_muc(doc, so, "Điểm mạnh và hạn chế")
    _hop(doc, "ĐIỂM NHẤN:", dn.get("manh_han_che"))
    _them(_doan(doc, sau=1), "ĐIỂM MẠNH", dam=True, mau=MAU_CHU["la"])
    _danh_sach(doc, tt.get("diem_manh"), mau_dau="la")
    _them(_doan(doc, truoc=4, sau=1), "HẠN CHẾ", dam=True, mau=MAU_CHU["do"])
    _danh_sach(doc, tt.get("han_che"), mau_dau="do")

    # Lưu ý CUỐI văn bản — nơi DUY NHẤT của «cần kiểm chứng», «cần xác nhận», nguồn gốc máy của GRADE (bác sĩ 19/10/2026).
    _doan(doc, sau=6)
    _them(_doan(doc, truoc=6, sau=4), "Lưu ý", dam=True, mau=MAU_CHU["cam"], co=13)
    _hop(doc, "CẦN BÁC SĨ KIỂM CHỨNG:", f"bản tóm tắt do máy (Claude) soạn từ toàn văn PDF. Mức chắc chắn GRADE ở mục "
         f"{so_grade} do máy đánh giá theo phương pháp GRADE; phần áp dụng vào thực hành là gợi ý — quyết định cho từng người "
         "bệnh thuộc bác sĩ (Cổng A).", mau="cam")
    if not _rong(tt.get("can_xac_nhan")):
        _them(_doan(doc, truoc=4, sau=1), "ĐIỂM CẦN XÁC NHẬN", dam=True, mau=MAU_CHU["cam"])
        _danh_sach(doc, tt.get("can_xac_nhan"), mau_dau="cam")
    _doan(doc, sau=4)
    _bang(doc, ["Nguồn của bản tóm tắt", ""], [
        ["Tệp toàn văn", f"{_o(pdf.get('tep'))} · {_o(pdf.get('so_trang'))} trang · SHA-256 {str(pdf.get('sha256', ''))[:16]}…"],
        ["Bản ghi EndNote", f"recNum {_o(bg.get('rec'))}"],
        ["Người soạn", tt.get("nguoi_soan") or "Claude (máy)"],
        ["Ngày dựng", dt.date.today().strftime("%d/%m/%Y")],
        ["Khuôn", f"phiên bản {PHIEN_BAN_KHUON} — tools/tom_tat_chung_cu_word.py"]], co=9.5, rong=[4.0, 13.0])
    _them(_doan(doc), "Không chứa thông tin định danh người bệnh.", nghieng=True, mau=MAU_CHU["xam"], co=10)
    return doc


def dung(thu_vien: Path, ma: str, *, gan_endnote: bool = False, ap_dung: bool = False, endnote=None,
         rut_bai=kiem_rut_bai) -> int:
    bg = tim_ban_ghi(thu_vien, ma)
    if bg is None:
        print(f"⚪ không thấy bản ghi có PDF cho «{ma}»")
        return 2
    goi = CACHE / bg["khoa"]
    if not (goi / "ngu_canh.json").is_file() or not (goi / "van_ban.txt").is_file():
        print(f"⚪ chưa có gói đọc — chạy `chuan-bi {bg['khoa']}` trước")
        return 2
    if not (goi / "tom_tat.json").is_file():
        print(f"⚪ chưa có {goi / 'tom_tat.json'} — phiên Claude đọc toàn văn và viết theo HUONG-DAN-DOC.md")
        return 2
    nc = json.loads((goi / "ngu_canh.json").read_text(encoding="utf-8"))
    if not (nc.get("dinh_danh") or {}).get("khop"):
        print(f"⛔ định danh PDF không khớp bản ghi ({(nc.get('dinh_danh') or {}).get('ly_do')}) — chặn")
        return 3
    nc["rut_bai"] = rut_bai(bg["pmid"], bg["doi"])  # tra LẠI lúc dựng — không dựa kết quả cũ của gói
    if nc["rut_bai"]["trang_thai"] == "da_rut":
        print("⛔ BÀI ĐÃ BỊ RÚT — không dựng tóm tắt để áp dụng")
        return 3
    try:
        tt = json.loads((goi / "tom_tat.json").read_text(encoding="utf-8"))
    except ValueError as e:
        print(f"✗ tom_tat.json không phải JSON hợp lệ: {e}")
        return 1
    loi = kiem_tom_tat(tt, doc_van_ban(goi), nc)
    if loi:
        print(f"✗ CỔNG: {len(loi)} lỗi — sửa tom_tat.json rồi chạy lại:")
        for x in loi:
            print(f"  · {x}")
        return 1
    loi = kiem_ket_qua_kiem_cheo(goi, tt)
    if loi:
        print(f"✗ KIỂM CHÉO ĐỘC LẬP: {len(loi)} lỗi — sửa tom_tat.json (nếu sai) rồi kiểm chéo lại:")
        for x in loi[:40]:
            print(f"  · {x}")
        return 1
    ra = tep_word(thu_vien, bg["khoa"])
    ra.parent.mkdir(parents=True, exist_ok=True)
    (ra.parent / "_nguon").mkdir(exist_ok=True)
    dung_word(tt, nc).save(str(ra))
    (ra.parent / "_nguon" / f"{bg['khoa']}.json").write_text(json.dumps(tt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"✓ Word: {ra}")
    pdf = dung_pdf(ra)
    print(f"✓ PDF: {pdf}" if pdf else "⚪ chưa có PDF — sẽ gắn bản Word (EndNote không xem trực tiếp được)")
    if gan_endnote:
        if pdf:
            return gan_vao_endnote(thu_vien, bg, pdf, ap_dung=ap_dung, endnote=endnote, thay=(ra.name,))
        return gan_vao_endnote(thu_vien, bg, ra, ap_dung=ap_dung, endnote=endnote)
    return 0


# ───────────────────────── gắn vào EndNote (THÊM, không thay) ─────────────────────────

def _tach_tep(gia_tri: str) -> list[str]:
    return [x.strip() for x in re.split(r"[\r\n]+", gia_tri or "") if x.strip()]


def _mo_endnote(thu_vien: Path):
    """EndNote đang mở ĐÚNG thư viện ⇒ đối tượng điều khiển; chưa mở ⇒ None (không tự bật EndNote)."""
    nec = mod("nec")
    ch = nec.doc_cau_hinh() or {}
    endnote = nec.EndNoteMac(ch.get("ung_dung") or "EndNote 2025", thu_vien)
    return endnote if endnote.dang_mo_dung_thu_vien() else None


def _ten_tep(x: str) -> str:
    """Tên tệp của một mục «File Attachments» (internal-pdf://…, file:///…, đường dẫn) — bỏ mã hoá %20."""
    return unquote(x.rstrip("/").rsplit("/", 1)[-1])


def gan_vao_endnote(thu_vien: Path, bg: dict, tep: Path, *, ap_dung: bool, endnote=None, thay: tuple[str, ...] = ()) -> int:
    """THÊM `tep` vào «File Attachments» của bản ghi; mục có tên trong `thay` (vd bản Word cũ cùng tên khi gắn PDF) được BỎ
    khỏi trường — tệp trên đĩa giữ nguyên. Mọi tệp còn lại (PDF toàn văn) phải còn nguyên khi đọc lại, thiếu ⇒ ghi trả."""
    if endnote is None:
        endnote = _mo_endnote(thu_vien)
        if endnote is None:
            print(f"⚪ EndNote chưa mở thư viện {thu_vien.name} — tệp đã lưu, chưa gắn")
            return 2
    ma = f'<RecordID database="{endnote.so_csdl()}" recNum="{bg["rec"]}" />'
    cu = _tach_tep(endnote.doc_truong([ma], "File Attachments").get(ma, ""))
    if not cu:
        print("⚪ không đọc được tệp đính kèm hiện có của bản ghi — không ghi (tránh xoá PDF)")
        return 2
    ten_thay = {t for t in thay if t != tep.name}
    giu = [x for x in cu if _ten_tep(x) not in ten_thay]
    bo = [x for x in cu if _ten_tep(x) in ten_thay]
    if any(_ten_tep(x) == tep.name for x in giu) and not bo:
        print(f"= đã gắn sẵn {tep.name} vào recNum {bg['rec']}")
        return 0
    giu_lai = [x for x in giu if _ten_tep(x) != tep.name]  # tệp cùng tên cũ (nếu có) thay bằng đường dẫn hiện hành
    moi = "\r".join(giu_lai + [tep.resolve().as_uri()])
    them_bo = f", bỏ liên kết {', '.join(_ten_tep(x) for x in bo)}" if bo else ""
    if not ap_dung:
        print(f"(chạy khô) sẽ thêm {tep.name} vào recNum {bg['rec']} — giữ {len(giu_lai)} tệp{them_bo}; thêm --ap-dung để ghi")
        return 0
    endnote._truoc_khi_ghi("gắn tóm tắt chứng cứ vào bản ghi")
    doc_lai = _tach_tep(endnote._chay('set field "File Attachments" of record (item 1 of argv) to (item 2 of argv)\n'
                                      'return field "File Attachments" of record (item 1 of argv)', ma, moi))
    if not all(x in doc_lai for x in giu_lai):
        endnote._chay('set field "File Attachments" of record (item 1 of argv) to (item 2 of argv)', ma, "\r".join(cu))
        print(f"✗ EndNote làm mất tệp cũ khi ghi — đã ghi trả lại {len(cu)} tệp cũ cho recNum {bg['rec']}")
        return 1
    if not any(_ten_tep(x) == tep.name for x in doc_lai):
        print(f"✗ EndNote không nhận {tep.name} (đọc lại không thấy)")
        return 1
    print(f"✓ đã gắn {tep.name} vào recNum {bg['rec']} (giữ {len(giu_lai)} tệp{them_bo})")
    return 0


def pdf_hoa(thu_vien: Path, *, ap_dung: bool, endnote=None, dung_pdf_fn=None, tim=None) -> int:
    """Mọi bản Word tóm tắt đã có → PDF giữ màu (dựng khi thiếu hoặc cũ hơn Word) rồi gắn PDF THAY liên kết Word trong
    EndNote. Chạy khô mặc định: không dựng tệp, không ghi EndNote. Mã: 0 xong · 1 có mục lỗi · 2 không đo được."""
    dung_pdf_fn = dung_pdf_fn or dung_pdf
    tim = tim or tim_ban_ghi
    ds = sorted(thu_muc_ra(thu_vien).glob("*_TomTatChungCu.docx"))
    if not ds:
        print(f"⚪ không có bản Word tóm tắt nào trong {thu_muc_ra(thu_vien)}")
        return 2
    if ap_dung and endnote is None:
        endnote = _mo_endnote(thu_vien)
        if endnote is None:
            print(f"⚪ EndNote chưa mở thư viện {thu_vien.name} — mở thư viện rồi chạy lại (không tự bật EndNote)")
            return 2
    loi = 0
    for docx in ds:
        khoa = docx.name[: -len("_TomTatChungCu.docx")]
        pdf = tep_pdf(docx)
        if not pdf_con_moi(docx):
            if not ap_dung:
                print(f"(chạy khô) {khoa}: sẽ dựng {pdf.name}")
            else:
                pdf = dung_pdf_fn(docx)
                if pdf is None:
                    loi += 1
                    continue
                print(f"✓ PDF: {pdf.name}")
        bg = tim(thu_vien, khoa)
        if bg is None:
            print(f"⚪ {khoa}: không thấy bản ghi có PDF trong EndNote — PDF đã có, chưa gắn")
            loi += 1
            continue
        if not ap_dung:
            print(f"(chạy khô) {khoa}: sẽ gắn {pdf.name} vào recNum {bg['rec']}, thay liên kết {docx.name}")
            continue
        if gan_vao_endnote(thu_vien, bg, pdf, ap_dung=True, endnote=endnote, thay=(docx.name,)) != 0:
            loi += 1
    print(f"{'(chạy khô) ' if not ap_dung else ''}{len(ds)} bản tóm tắt · {loi} mục lỗi")
    return 1 if loi else 0


# ───────────────────────── MỐC «để nguyên» + HÀNG CHỜ chứng cứ MỚI ─────────────────────────
# Bác sĩ 19/10/2026: «Bây giờ xây dựng cho những chứng cứ sau này, còn hiện tại hãy để nguyên». Mốc liệt kê MỌI bản ghi đang
# có (kể cả trong thùng rác) lúc ghi — các bản ghi đó KHÔNG BAO GIỜ vào hàng chờ, kể cả khi sau này được gắn thêm PDF. Chỉ bản
# ghi nạp SAU mốc (recNum mới) có PDF và chưa có Word mới vào hàng chờ. Thiếu mốc ⇒ KHÔNG đoán «mới» là gì (fail-closed).

TEN_MOC = "_moc_de_nguyen.json"
LOI_MOC_TOI_THIEU = 10


def tep_moc(thu_vien: Path) -> Path:
    return thu_muc_ra(thu_vien) / TEN_MOC


def ghi_moc(thu_vien: Path, loi_bac_si: str, *, ghi_de: bool = False) -> int:
    tep = tep_moc(thu_vien)
    if tep.exists() and not ghi_de:
        print(f"= đã có mốc {tep} — không ghi đè (dùng --ghi-de nếu bác sĩ muốn dời mốc)")
        return 1
    if len((loi_bac_si or "").strip()) < LOI_MOC_TOI_THIEU:
        print("✗ cần NGUYÊN VĂN lời bác sĩ (--loi) làm căn cứ ghi mốc")
        return 1
    c = mod("sk")._ket_noi(thu_vien)
    try:
        ids = sorted(r[0] for r in c.execute("select id from refs"))
    finally:
        c.close()
    tep.parent.mkdir(parents=True, exist_ok=True)
    tep.write_text(json.dumps({"phien_ban": 1, "luc": dt.datetime.now().isoformat(timespec="seconds"),
                               "loi_bac_si": loi_bac_si.strip(), "thu_vien": thu_vien.name, "so_ban_ghi": len(ids),
                               "rec": ids}, ensure_ascii=False), encoding="utf-8")
    print(f"✓ mốc «để nguyên»: {len(ids)} bản ghi hiện có KHÔNG vào hàng chờ tóm tắt — {tep}")
    return 0


def doc_moc(thu_vien: Path) -> set[int] | None:
    """Tập recNum thuộc mốc; None = chưa có mốc hoặc mốc hỏng (KHÔNG coi là rỗng — rỗng sẽ đẩy cả thư viện vào hàng chờ)."""
    try:
        d = json.loads(tep_moc(thu_vien).read_text(encoding="utf-8"))
        ids = d["rec"]
        if not isinstance(ids, list) or not ids:
            return None
        return {int(x) for x in ids}
    except (OSError, ValueError, KeyError, TypeError):
        return None


# Bác sĩ 24/10/2026: «Hãy kiểm tra lại tất cả các trích dẫn trong Endnote giúp tôi,đảm bảo đã được tải,đọc toàn văn và tóm tắt
# chứng cứ một cách chi tiết» → chọn phương án «T1 + T2». NGOẠI LỆ của mốc theo TẦNG ưu tiên: bản ghi trong mốc mà là nguồn của mục
# «Áp dụng» (T1) / «Cân nhắc» (T2) vẫn vào hàng chờ; T3/T4 giữ luật mốc. Tệp vắng/hỏng ⇒ không ngoại lệ (đúng luật 19/10).
TEN_NGOAI_LE = "_ngoai_le_moc.json"


class NgoaiLeKhongDoDuoc(RuntimeError):
    """Có ngoại lệ mốc nhưng không đọc được dashboard nào để xếp tầng (vd chạy ở worktree/bản sao trần) — KHÔNG được trả «0 chờ»."""


def tep_ngoai_le(thu_vien: Path) -> Path:
    return thu_muc_ra(thu_vien) / TEN_NGOAI_LE


def ghi_ngoai_le_moc(thu_vien: Path, tang: list[int], loi_bac_si: str) -> int:
    """Ghi quyết định ngoại lệ mốc (thay quyết định trước). Bắt buộc nguyên văn lời bác sĩ và tầng thuộc 1–4."""
    if len((loi_bac_si or "").strip()) < LOI_MOC_TOI_THIEU:
        print("✗ cần NGUYÊN VĂN lời bác sĩ (--loi) làm căn cứ ghi ngoại lệ mốc")
        return 1
    tap = sorted({int(t) for t in tang})
    if not tap or any(t not in (1, 2, 3, 4) for t in tap):
        print("✗ --tang phải là một hay nhiều giá trị trong 1 2 3 4")
        return 1
    tep = tep_ngoai_le(thu_vien)
    tep.parent.mkdir(parents=True, exist_ok=True)
    tep.write_text(json.dumps({"phien_ban": 1, "luc": dt.datetime.now().isoformat(timespec="seconds"),
                               "loi_bac_si": loi_bac_si.strip(), "tang": tap}, ensure_ascii=False), encoding="utf-8")
    print(f"✓ ngoại lệ mốc: bản ghi cũ tầng {', '.join(f'T{t}' for t in tap)} vào hàng chờ tóm tắt — {tep}")
    return 0


def doc_ngoai_le_moc(thu_vien: Path) -> set[int]:
    """Tầng được miễn mốc; vắng/hỏng/sai kiểu ⇒ tập rỗng (giữ nguyên luật mốc 19/10)."""
    try:
        d = json.loads(tep_ngoai_le(thu_vien).read_text(encoding="utf-8"))
        tang = d["tang"]
        if not isinstance(tang, list) or len(str(d.get("loi_bac_si") or "").strip()) < LOI_MOC_TOI_THIEU:
            return set()
        return {int(t) for t in tang if int(t) in (1, 2, 3, 4)}
    except (OSError, ValueError, KeyError, TypeError):
        return set()


def trich_dan_cua_dashboard(dash: Path) -> tuple[set[str], set[str]]:
    """(PMID, DOI chuẩn hoá) mà dashboard trích — mục chính lẫn tài liệu tham khảo."""
    ds = mod("tt").ngu_canh([dash], mod("nec")._nap_doc_data())
    pm = {v["pmid"] for v in ds.values() if v.get("pmid")}
    doi = {v["doi"] for v in ds.values() if v.get("doi")}
    return pm, doi


# Bác sĩ 25/10/2026: trích dẫn của ĐỀ TÀI nghiên cứu (nhóm «NC · …», dựng bởi `nap_endnote_chu_de.py --de-tai`) không cần tóm tắt
# chi tiết, nhưng phải có toàn văn ⇒ bản ghi CHỈ thuộc nhóm «NC · » (không nằm trong nhóm «EBM · » nào) không vào hàng chờ. Bản ghi
# vừa là trích dẫn đề tài vừa nằm trong nhóm chứng cứ vẫn theo luật chứng cứ. Đo 25/10: hàng chờ nhảy 2 → 43, 41 bài chỉ thuộc nhóm
# «NC · Chăm sóc thận ĐTĐ típ 2…». Nhóm đo theo group_id trên bản sao sdb.eni (BH167), không qua AppleScript.

#
# 26/10/2026: bác sĩ cho loại 19/66 trích dẫn của đề tài btm-dtd2 (không toàn văn hoặc ít phù hợp) ⇒ phiên đề tài gỡ chúng khỏi nhóm
# «NC · » (không thùng rác). 14 bài có PDF trong số đó không còn nhóm nào nên rơi vào hàng chờ như «chứng cứ MỚI» (đo: 16 bài chờ, 14
# là bài đề tài đã loại). Nhóm thôi là chưa đủ: bài có trong danh mục trích dẫn của MỘT đề tài (`trich_dan` lẫn `loai_bo` của
# `DANH_MUC_TRICH_DAN_<mã>.json`) mà không thuộc nhóm «EBM · » nào cũng không vào hàng chờ.

THU_MUC_DE_TAI: Path | None = None  # None ⇒ <repo y khoa>/exports qua goc_du_lieu (lồng rồi anh em); test gán thư mục tạm


class NhomKhongDoDuoc(RuntimeError):
    """Không đọc được nhóm EndNote để tách trích dẫn đề tài — KHÔNG được trả hàng chờ lẫn bài đề tài như thể đã lọc."""


def trich_dan_danh_muc_de_tai(thu_muc: Path | None = None) -> tuple[set[str], set[str]]:
    """(PMID, DOI chữ thường) trong `trich_dan` + `loai_bo` của mọi `<đề tài>/DANH_MUC_TRICH_DAN_*.json` (bỏ bản `.bak`). Vắng thư
    mục đề tài (phiên Cloud, worktree) ⇒ hai tập rỗng; danh mục không đọc được hoặc sai khuôn ⇒ NhomKhongDoDuoc (không đoán)."""
    goc = thu_muc or THU_MUC_DE_TAI or (goc_du_lieu("medical-ebm-automation") / "exports")
    pm: set[str] = set()
    doi: set[str] = set()
    if not goc.is_dir():
        return pm, doi
    for tep in sorted(goc.glob("*/DANH_MUC_TRICH_DAN_*.json")):
        if ".bak" in tep.name:
            continue
        try:
            d = json.loads(tep.read_text(encoding="utf-8"))
            muc = list(d.get("trich_dan") or []) + list(d.get("loai_bo") or [])
            for m in muc:
                if m.get("pmid"):
                    pm.add(str(m["pmid"]).strip())
                if m.get("doi"):
                    doi.add(str(m["doi"]).strip().lower())
        except (OSError, ValueError, AttributeError, TypeError) as e:
            raise NhomKhongDoDuoc(f"danh mục trích dẫn đề tài {tep.parent.name}/{tep.name} không đọc được — "
                                  f"{type(e).__name__}: {e}"[:300]) from e
    return pm, doi


def _nhom_nc_ebm(thu_vien: Path) -> tuple[set[int], set[int]]:
    """(recNum trong nhóm «NC · », recNum trong nhóm «EBM · »). Không đọc được bảng nhóm, hoặc thành viên của một nhóm «NC · »/
    «EBM · » không giải được ⇒ NhomKhongDoDuoc (không đoán)."""
    sk, nec = mod("sk"), mod("nec")
    tien_to_nc, tien_to_ebm = sk.khoa_ten(nec.TIEN_TO_NHOM_DE_TAI), sk.khoa_ten(nec.TIEN_TO_NHOM)
    try:
        nhom = sk.doc_nhom(thu_vien)
    except (OSError, sqlite3.Error) as e:
        raise NhomKhongDoDuoc(f"không đọc được nhóm EndNote để tách trích dẫn đề tài «{nec.TIEN_TO_NHOM_DE_TAI}…» — "
                              f"{type(e).__name__}: {e}"[:300]) from e
    nc: set[int] = set()
    ebm: set[int] = set()
    for n in nhom:
        ten = sk.khoa_ten(n.get("ten") or "")
        dich = nc if ten.startswith(tien_to_nc) else ebm if ten.startswith(tien_to_ebm) else None
        if dich is None:
            continue
        if n.get("rec") is None:
            raise NhomKhongDoDuoc(f"nhóm {n.get('group_id')} «{n.get('ten')}»: {n.get('loi') or 'không giải được thành viên'}")
        dich.update(n["rec"])
    return nc, ebm


def chi_thuoc_de_tai(thu_vien: Path) -> set[int]:
    """recNum nằm trong ≥ 1 nhóm «NC · » và không nằm trong nhóm «EBM · » nào (lỗi nhóm ⇒ NhomKhongDoDuoc)."""
    nc, ebm = _nhom_nc_ebm(thu_vien)
    return nc - ebm


def hang_cho(thu_vien: Path, *, dashboard: Path | None = None) -> list[dict] | None:
    """Trích dẫn MỚI (ngoài mốc) có PDF và chưa có Word, rồi bản ghi CŨ thuộc tầng ngoại lệ mốc (`ngoai-le-moc`, xếp theo tầng).
    Mỗi mục mang `nguon_hang_cho`: «moi» hoặc «ngoai_le_moc». None = chưa có mốc (không đoán). Bản ghi KHÔNG thuộc nhóm «EBM · »
    nào mà thuộc nhóm đề tài «NC · » hoặc có PMID/DOI trong danh mục trích dẫn của một đề tài (kể cả bài đã loại) bị loại; không đọc
    được nhóm hay danh mục ⇒ NhomKhongDoDuoc."""
    moc = doc_moc(thu_vien)
    if moc is None:
        return None
    co_pdf = [bg for bg in ban_ghi_co_pdf(thu_vien) if not tep_word(thu_vien, bg["khoa"]).is_file()]
    moi = [dict(bg, nguon_hang_cho="moi") for bg in co_pdf if bg["rec"] not in moc]
    cu: list[dict] = []
    ngoai_le = doc_ngoai_le_moc(thu_vien)
    if ngoai_le:
        ds = ngu_canh_dashboard()
        if not ds:
            raise NgoaiLeKhongDoDuoc("không đọc được dashboard nào (EBM-Dashboards/) để xếp tầng bản ghi cũ — không đo được")
        for bg in co_pdf:
            if bg["rec"] in moc:
                t = tang_uu_tien(_ngu_canh_cua(ds, bg["pmid"], bg["doi"]))
                if t in ngoai_le:
                    cu.append(dict(bg, tang=t, nguon_hang_cho="ngoai_le_moc"))
    if moi or cu:  # chỉ đọc nhóm khi có ứng viên (đọc nhóm = chép tạm sdb.eni)
        nc, ebm = _nhom_nc_ebm(thu_vien)
        pm_dt, doi_dt = trich_dan_danh_muc_de_tai()

        def la_de_tai(bg: dict) -> bool:
            if bg["rec"] in ebm:  # vừa là trích dẫn đề tài vừa là chứng cứ ⇒ theo luật chứng cứ
                return False
            return (bg["rec"] in nc or (bool(bg.get("pmid")) and str(bg["pmid"]) in pm_dt)
                    or (bool(bg.get("doi")) and str(bg["doi"]).lower() in doi_dt))

        moi, cu = ([bg for bg in nh if not la_de_tai(bg)] for nh in (moi, cu))
    if dashboard is not None:
        pm, doi = trich_dan_cua_dashboard(dashboard)
        moi, cu = ([bg for bg in nh if (bg["pmid"] and bg["pmid"] in pm) or (bg["doi"] and bg["doi"] in doi)] for nh in (moi, cu))
    return sorted(moi, key=lambda x: x["rec"]) + sorted(cu, key=lambda x: (x["tang"], x["rec"]))


HUONG_DAN_HANG_CHO = (
    "Mỗi trích dẫn trong hàng chờ (phiên Claude, TUẦN TỰ — mỗi bài ≈ 150.000–400.000 token):\n"
    "  1. `chuan-bi <mã>` (gói đọc + tra rút bài) — bài rút / PDF sai bài ⇒ dừng, báo bác sĩ.\n"
    "  2. MỘT agent đọc toàn văn theo HUONG-DAN-DOC.md → tom_tat.json → `kiem <mã>` tới khi đạt.\n"
    "  3. `kiem-cheo <mã>` → MỘT agent ĐỘC LẬP khác chấm từng câu → sửa mục sai (đối chiếu lại bài) → lặp tới 0 sai.\n"
    "  4. `dung <mã> --gan-endnote --ap-dung` (Word cạnh thư viện, gắn liên kết, giữ PDF).")


# ───────────────────────── liệt kê / tiến độ ─────────────────────────

def liet_ke(thu_vien: Path, *, tang: int | None = None, chua_lam: bool = False) -> list[dict]:
    ds = ngu_canh_dashboard()
    ra = []
    for bg in ban_ghi_co_pdf(thu_vien):
        nc = _ngu_canh_cua(ds, bg["pmid"], bg["doi"])
        bg["tang"] = tang_uu_tien(nc)
        bg["da_lam"] = tep_word(thu_vien, bg["khoa"]).is_file()
        bg["muc_dashboard"] = sorted({f"{x.get('muc')}:{x.get('decision')}" for x in nc if x.get("vai") == "muc"})
        if (tang is None or bg["tang"] == tang) and not (chua_lam and bg["da_lam"]):
            ra.append(bg)
    return sorted(ra, key=lambda x: (x["tang"], x["da_lam"], x["khoa"]))


def _thu_vien() -> Path | None:
    tv = mod("sk").thu_vien_cau_hinh()
    if tv is None or not Path(tv).is_file():
        print("⚪ chưa khai thư viện EndNote (~/.ebm-endnote.json) hoặc tệp .enl vắng")
        return None
    return Path(tv)


def main(argv: list[str] | None = None) -> int:
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description="Tóm tắt chứng cứ chi tiết tiếng Việt → Word Times New Roman, gắn EndNote")
    sub = ap.add_subparsers(dest="lenh", required=True)
    p = sub.add_parser("liet-ke")
    p.add_argument("--tang", type=int, choices=(1, 2, 3, 4))
    p.add_argument("--chua-lam", action="store_true")
    p.add_argument("--json", action="store_true")
    sub.add_parser("mau")
    sub.add_parser("tien-do")
    p = sub.add_parser("moc", help="ghi mốc «để nguyên»: bản ghi hiện có không vào hàng chờ tóm tắt")
    p.add_argument("--loi", required=True, help="NGUYÊN VĂN lời bác sĩ")
    p.add_argument("--ghi-de", action="store_true")
    p = sub.add_parser("ngoai-le-moc", help="bản ghi CŨ (trong mốc) thuộc tầng ưu tiên đã chọn vẫn vào hàng chờ tóm tắt")
    p.add_argument("--tang", type=int, nargs="+", required=True, choices=(1, 2, 3, 4))
    p.add_argument("--loi", required=True, help="NGUYÊN VĂN lời bác sĩ")
    p = sub.add_parser("pdf-hoa", help="mọi Word đã có → PDF giữ màu, gắn PDF thay Word trong EndNote")
    p.add_argument("--ap-dung", action="store_true")
    p = sub.add_parser("hang-cho", help="trích dẫn MỚI (ngoài mốc) có PDF chưa có Word")
    p.add_argument("--dashboard", type=Path)
    p.add_argument("--chuan-bi", action="store_true", help="dựng sẵn gói đọc (tra rút bài, cần mạng) cho từng mục")
    p.add_argument("--json", action="store_true")
    for ten in ("chuan-bi", "kiem", "kiem-cheo", "dung"):
        p = sub.add_parser(ten)
        p.add_argument("ma")
        if ten == "dung":
            p.add_argument("--gan-endnote", action="store_true")
            p.add_argument("--ap-dung", action="store_true")
    a = ap.parse_args(argv)
    if a.lenh == "mau":
        print(huong_dan_doc())
        print(json.dumps(MAU, ensure_ascii=False, indent=2))
        return 0
    tv = _thu_vien()
    if tv is None:
        return 2
    if a.lenh == "liet-ke":
        ds = liet_ke(tv, tang=a.tang, chua_lam=a.chua_lam)
        if a.json:
            print(json.dumps([{k: v for k, v in x.items() if k != "pdf"} for x in ds], ensure_ascii=False, indent=1))
        else:
            for x in ds:
                print(f"T{x['tang']} {'✓' if x['da_lam'] else '·'} {x['khoa']:<14} rec {x['rec']:<5} {x['tieu_de'][:70]}"
                      f"  {' '.join(x['muc_dashboard'])[:40]}")
            print(f"— {len(ds)} trích dẫn")
        return 0
    if a.lenh == "moc":
        return ghi_moc(tv, a.loi, ghi_de=a.ghi_de)
    if a.lenh == "ngoai-le-moc":
        return ghi_ngoai_le_moc(tv, a.tang, a.loi)
    if a.lenh == "pdf-hoa":
        return pdf_hoa(tv, ap_dung=a.ap_dung)
    if a.lenh == "hang-cho":
        try:
            hc = hang_cho(tv, dashboard=a.dashboard)
        except (NgoaiLeKhongDoDuoc, NhomKhongDoDuoc) as e:
            print(f"⚪ {e}")
            return 2
        if hc is None:
            print(f"⚪ chưa có mốc «để nguyên» ({tep_moc(tv)}) — không xác định được trích dẫn MỚI; bác sĩ quyết rồi chạy "
                  "`moc --loi \"<nguyên văn>\"`")
            return 2
        if a.json:
            print(json.dumps([{k: v for k, v in x.items() if k != "pdf"} for x in hc], ensure_ascii=False, indent=1))
            return 0
        if not hc:
            print("🟢 không có trích dẫn MỚI có PDF chờ tóm tắt Word" + (f" trong {a.dashboard.name}" if a.dashboard else ""))
            return 0
        print(f"🛎 {len(hc)} trích dẫn MỚI có PDF chờ tóm tắt Word" + (f" ({a.dashboard.name})" if a.dashboard else "") + ":")
        for x in hc:
            print(f"  · {x['khoa']:<14} rec {x['rec']:<5} {x['tieu_de'][:80]}")
            if a.chuan_bi:
                chuan_bi(tv, x["khoa"])
        print(HUONG_DAN_HANG_CHO)
        return 0
    if a.lenh == "tien-do":
        ds = liet_ke(tv)
        for t in (1, 2, 3, 4):
            nhom = [x for x in ds if x["tang"] == t]
            print(f"Tầng {t}: {sum(x['da_lam'] for x in nhom)}/{len(nhom)} đã có Word")
        print(f"Tổng: {sum(x['da_lam'] for x in ds)}/{len(ds)} · thư mục {thu_muc_ra(tv)}")
        word = sorted(thu_muc_ra(tv).glob("*_TomTatChungCu.docx"))
        print(f"PDF giữ màu (mở được trong EndNote): {sum(pdf_con_moi(w) for w in word)}/{len(word)}"
              + ("" if all(pdf_con_moi(w) for w in word) else " — chạy `pdf-hoa --ap-dung`"))
        try:
            hc = hang_cho(tv)
        except (NgoaiLeKhongDoDuoc, NhomKhongDoDuoc) as e:
            print(f"Hàng chờ: ⚪ {e}")
            return 2
        if hc is None:
            print("Hàng chờ chứng cứ MỚI: ⚪ chưa có mốc «để nguyên»")
        else:
            print(f"Hàng chờ chứng cứ MỚI: {sum(x['nguon_hang_cho'] == 'moi' for x in hc)} trích dẫn")
            nl = doc_ngoai_le_moc(tv)
            if nl:
                cu = [x for x in hc if x["nguon_hang_cho"] == "ngoai_le_moc"]
                print(f"Ngoại lệ mốc ({', '.join(f'T{t}' for t in sorted(nl))}): còn {len(cu)} bản ghi cũ chờ tóm tắt — "
                      + " · ".join(f"T{t} {sum(x['tang'] == t for x in cu)}" for t in sorted(nl)))
        return 0
    if a.lenh == "chuan-bi":
        return chuan_bi(tv, a.ma)[0]
    if a.lenh == "kiem-cheo":
        bg = tim_ban_ghi(tv, a.ma)
        goi = CACHE / bg["khoa"] if bg else None
        if not goi or not (goi / "tom_tat.json").is_file():
            print("⚪ chưa có gói đọc/tom_tat.json")
            return 2
        loi = kiem_tom_tat(json.loads((goi / "tom_tat.json").read_text(encoding="utf-8")), doc_van_ban(goi),
                           json.loads((goi / "ngu_canh.json").read_text(encoding="utf-8")))
        if loi:
            print(f"✗ cổng tất định còn {len(loi)} lỗi — sửa trước khi kiểm chéo (chạy `kiem {bg['khoa']}`)")
            return 1
        ra = chuan_bi_kiem_cheo(goi)
        print(f"Gói kiểm chéo: {ra} ({len(json.loads(ra.read_text(encoding='utf-8'))['muc'])} mục)")
        print(f"  Giao một agent ĐỘC LẬP (ngữ cảnh mới) đọc {goi / 'HUONG-DAN-KIEM-CHEO.md'} → viết {goi / 'kiem_cheo.json'}")
        return 0
    if a.lenh == "kiem":
        bg = tim_ban_ghi(tv, a.ma)
        goi = CACHE / bg["khoa"] if bg else None
        if not goi or not (goi / "tom_tat.json").is_file():
            print("⚪ chưa có gói đọc/tom_tat.json")
            return 2
        loi = kiem_tom_tat(json.loads((goi / "tom_tat.json").read_text(encoding="utf-8")), doc_van_ban(goi),
                           json.loads((goi / "ngu_canh.json").read_text(encoding="utf-8")))
        for x in loi:
            print(f"  · {x}")
        print("✓ cổng đạt" if not loi else f"✗ {len(loi)} lỗi")
        return 1 if loi else 0
    return dung(tv, a.ma, gan_endnote=a.gan_endnote, ap_dung=a.ap_dung)


if __name__ == "__main__":
    raise SystemExit(main())
