from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import batch_evidence_quality_check as Q  # noqa: E402


def _card(card_id: str, *, decision: str = "notyet", pmid: str = "12345678") -> dict:
    return {
        "id": card_id,
        "topic": "Hypertension treatment safety monitoring",
        "specialty": "cardiology",
        "recommendation": "Treat hypertension and monitor adverse events, bleeding risk, renal safety.",
        "date_source": "2025",
        "gradeLevel": "high",
        "decision": decision,
        "verification_status": "da xac minh",
        "certainty": "high certainty with safety discussion",
        "source": {
            "pmid": pmid,
            "doi": "10.1000/example",
            "agency": "WHO",
            "title": "WHO guideline on hypertension treatment and safety",
        },
        "critical_appraisal": {
            "design": "guideline",
            "effect_estimate": "benefit with adverse event monitoring",
        },
    }


def test_pending_scope_limits_to_78_notyet_cards():
    cards = [_card(f"notyet-{idx}", decision="notyet") for idx in range(100)]
    cards += [_card(f"consider-{idx}", decision="consider") for idx in range(25)]

    selected = Q.filter_cards_for_review(cards, scope="pending", limit=78)

    assert len(selected) == 78
    assert {card["decision"] for card in selected} == {"notyet"}


def test_auto_classification_is_review_queue_not_immediate_apply():
    checker = Q.EvidenceQualityChecker()

    results = checker.run_all_checks(_card("ready-1"))
    classification = checker.get_classification(results)

    assert classification == Q.AUTO_APPLY_REVIEW
    assert "APPLY_IMMEDIATELY" not in classification


def test_q2_failure_blocks_ready_for_apply_review():
    checker = Q.EvidenceQualityChecker()
    bad_source_card = _card("bad-source", pmid="")
    bad_source_card["source"]["doi"] = ""

    results = checker.run_all_checks(bad_source_card)

    assert results["Q2"][0] is False
    assert checker.get_classification(results) == Q.AUTO_NOT_READY


def test_export_workbook_has_doctor_decision_validation(tmp_path):
    output = tmp_path / "approval.xlsx"
    config = Q.WorkflowConfig(output=output, json_report=tmp_path / "approval.json")
    results = Q.run_quality_check_batch([_card("ready-1")])

    exported = Q.export_to_excel(results, config)

    workbook = openpyxl.load_workbook(exported)
    assert {"Evidence Approval", "README", "Q1-Q7 Criteria"}.issubset(workbook.sheetnames)

    worksheet = workbook["Evidence Approval"]
    headers = [cell.value for cell in worksheet[1]]
    assert worksheet.max_row == 2
    assert "Doctor\nDecision" in headers
    assert "Auto\nClassify" in headers
    assert any(
        validation.type == "list" and "APPLY" in str(validation.formula1)
        for validation in worksheet.data_validations.dataValidation
    )
