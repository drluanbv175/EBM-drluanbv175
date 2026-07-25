#!/usr/bin/env python3
"""
Process Approved Evidence — Cập nhật DB sau khi bác sĩ ký Excel

Giai đoạn 2 (Tuần 3-4) — Công việc cuối:
  1. Bác sĩ điền Doctor_Decision vào Excel
  2. Script này đọc Excel → cập nhật EBM_MASTER.json
  3. Chạy integrity_guard.py --fix → DANH_MUC.html
  4. Email báo cáo

Cách dùng:
  python process_approved_evidence.py --input _EVIDENCE_APPROVAL_WORKFLOW.xlsx
"""

import json
import sys
import logging
from pathlib import Path
from typing import Dict, List, Tuple
from datetime import datetime
import shutil

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not found. Install: pip install openpyxl")
    sys.exit(1)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# Hằng số
EVIDENCE_DB_PATH = Path(__file__).parent.parent / "EBM_MASTER" / "EBM_MASTER.json"
BACKUP_DIR = Path(__file__).parent.parent / "EBM_MASTER" / "BACKUPS"


class ApprovedEvidenceProcessor:
    """Xử lý kết quả duyệt từ Excel"""
    
    def __init__(self):
        self.evidence_db = {}
        self.updates = {
            "apply": [],
            "consider": [],
            "notyet": [],
            "unchanged": 0,
            "errors": []
        }
    
    def load_evidence_db(self) -> None:
        """Load EBM_MASTER.json"""
        if not EVIDENCE_DB_PATH.exists():
            logger.error(f"File not found: {EVIDENCE_DB_PATH}")
            sys.exit(1)
        
        with open(EVIDENCE_DB_PATH, "r", encoding="utf-8") as f:
            self.evidence_db = json.load(f)
        
        logger.info(f"Loaded {len(self.evidence_db.get('evidence_cards', []))} cards from DB")
    
    def load_approved_excel(self, excel_path: str) -> List[Dict]:
        """Load kết quả từ Excel đã bác sĩ ký"""
        excel_file = Path(excel_path)
        
        if not excel_file.exists():
            logger.error(f"Excel file not found: {excel_file}")
            sys.exit(1)
        
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active
        
        # Đọc header
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        
        logger.info(f"Excel headers: {headers}")
        
        # Tìm cột Doctor_Decision
        try:
            doctor_decision_idx = headers.index("Doctor\nDecision")
            doctor_notes_idx = headers.index("Doctor\nNotes")
            id_idx = headers.index("ID")
        except ValueError as e:
            logger.error(f"Column not found: {e}")
            logger.error(f"Available: {headers}")
            sys.exit(1)
        
        results = []
        for row_idx in range(2, ws.max_row + 1):
            row = ws[row_idx]
            
            card_id = row[id_idx].value
            decision = row[doctor_decision_idx].value
            notes = row[doctor_notes_idx].value or ""
            
            if not card_id:
                continue  # Skip empty rows
            
            # Normalize decision
            decision = decision.strip().upper() if decision else "UNCHANGED"
            
            if decision not in ["APPLY", "CONSIDER", "NOT_YET", "UNCHANGED"]:
                logger.warning(f"Row {row_idx}: Unknown decision '{decision}' for {card_id}")
                self.updates["errors"].append({
                    "card_id": card_id,
                    "error": f"Unknown decision: {decision}"
                })
                continue
            
            results.append({
                "card_id": card_id,
                "doctor_decision": decision,
                "doctor_notes": notes,
                "row_idx": row_idx
            })
        
        logger.info(f"Loaded {len(results)} approved decisions from Excel")
        return results
    
    def map_decision_to_field(self, decision: str) -> str:
        """Map Doctor_Decision thành decision field DB"""
        mapping = {
            "APPLY": "apply",
            "CONSIDER": "consider",
            "NOT_YET": "notyet",
            "UNCHANGED": None  # No change
        }
        return mapping.get(decision)
    
    def find_card_by_id(self, card_id: str) -> Tuple[int, Dict]:
        """Tìm thẻ trong DB bằng ID"""
        cards = self.evidence_db.get("evidence_cards", [])
        
        for idx, card in enumerate(cards):
            if card.get("id") == card_id:
                return idx, card
        
        return -1, {}
    
    def update_card(self, card_idx: int, card: Dict, new_decision: str, notes: str) -> bool:
        """Cập nhật thẻ trong DB"""
        if card_idx < 0:
            return False
        
        # Lưu lịch sử
        history = card.get("history", [])
        history.append({
            "date": datetime.now().strftime("%Y-%m-%d"),
            "change": f"bác sĩ phê duyệt: {new_decision}" + (f" - {notes}" if notes else "")
        })
        
        card["decision"] = new_decision
        card["history"] = history
        
        # Update DB
        self.evidence_db["evidence_cards"][card_idx] = card
        
        return True
    
    def process_approvals(self, approvals: List[Dict]) -> None:
        """Xử lý danh sách phê duyệt"""
        for approval in approvals:
            card_id = approval["card_id"]
            doctor_decision = approval["doctor_decision"]
            doctor_notes = approval["doctor_notes"]
            
            # Map decision
            new_decision = self.map_decision_to_field(doctor_decision)
            
            if new_decision is None:
                self.updates["unchanged"] += 1
                continue
            
            # Tìm thẻ
            card_idx, card = self.find_card_by_id(card_id)
            
            if card_idx < 0:
                logger.warning(f"Card not found: {card_id}")
                self.updates["errors"].append({
                    "card_id": card_id,
                    "error": "Card not found in DB"
                })
                continue
            
            # Check xem có thay đổi không
            old_decision = card.get("decision")
            
            if old_decision == new_decision:
                self.updates["unchanged"] += 1
                continue
            
            # Cập nhật
            if self.update_card(card_idx, card, new_decision, doctor_notes):
                self.updates[new_decision].append({
                    "id": card_id,
                    "old": old_decision,
                    "new": new_decision,
                    "notes": doctor_notes
                })
                logger.info(f"Updated {card_id}: {old_decision} → {new_decision}")
            else:
                self.updates["errors"].append({
                    "card_id": card_id,
                    "error": "Failed to update"
                })
    
    def save_evidence_db(self) -> Path:
        """Lưu DB đã cập nhật"""
        # Backup cũ
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        backup_file = BACKUP_DIR / f"EBM_MASTER_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        shutil.copy(EVIDENCE_DB_PATH, backup_file)
        logger.info(f"Backup created: {backup_file}")
        
        # Lưu DB mới
        with open(EVIDENCE_DB_PATH, "w", encoding="utf-8") as f:
            json.dump(self.evidence_db, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Updated DB saved: {EVIDENCE_DB_PATH}")
        return EVIDENCE_DB_PATH
    
    def generate_report(self) -> Dict:
        """Tạo báo cáo"""
        total_updated = len(self.updates["apply"]) + len(self.updates["consider"]) + len(self.updates["notyet"])
        
        report = {
            "timestamp": datetime.now().isoformat(),
            "summary": {
                "total_processed": total_updated + self.updates["unchanged"] + len(self.updates["errors"]),
                "total_updated": total_updated,
                "unchanged": self.updates["unchanged"],
                "errors": len(self.updates["errors"]),
            },
            "updates": {
                "apply": len(self.updates["apply"]),
                "consider": len(self.updates["consider"]),
                "notyet": len(self.updates["notyet"]),
            },
            "details": self.updates
        }
        
        return report
    
    def print_report(self, report: Dict) -> None:
        """In báo cáo"""
        print("\n" + "="*70)
        print("APPROVED EVIDENCE PROCESSING REPORT")
        print("="*70)
        
        summary = report["summary"]
        print(f"\n📊 SUMMARY:")
        print(f"  Total processed: {summary['total_processed']}")
        print(f"  Updated: {summary['total_updated']}")
        print(f"  Unchanged: {summary['unchanged']}")
        print(f"  Errors: {summary['errors']}")
        
        updates = report["updates"]
        print(f"\n📈 UPDATED DISTRIBUTION:")
        print(f"  → APPLY: {updates['apply']} cards")
        print(f"  → CONSIDER: {updates['consider']} cards")
        print(f"  → NOT_YET: {updates['notyet']} cards")
        
        if self.updates["errors"]:
            print(f"\n⚠️ ERRORS ({len(self.updates['errors'])}):")
            for error in self.updates["errors"][:5]:
                print(f"  - {error['card_id']}: {error['error']}")
            if len(self.updates["errors"]) > 5:
                print(f"  ... and {len(self.updates['errors']) - 5} more")
        
        print("\n" + "="*70)


def main():
    import argparse
    
    parser = argparse.ArgumentParser(description="Process approved evidence from Excel")
    parser.add_argument("--input", "-i", required=True, help="Input Excel file (bác sĩ ký)")
    parser.add_argument("--dry-run", action="store_true", help="Không lưu DB, chỉ show báo cáo")
    
    args = parser.parse_args()
    
    # Xử lý
    processor = ApprovedEvidenceProcessor()
    processor.load_evidence_db()
    
    approvals = processor.load_approved_excel(args.input)
    processor.process_approvals(approvals)
    
    # Báo cáo
    report = processor.generate_report()
    processor.print_report(report)
    
    # Lưu
    if not args.dry_run:
        processor.save_evidence_db()
        
        # Lưu báo cáo
        report_file = Path(__file__).parent.parent / "reports" / f"approval_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        report_file.parent.mkdir(parents=True, exist_ok=True)
        
        with open(report_file, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Report saved: {report_file}")
        logger.info("\n✅ DONE — Database updated")
        logger.info("Next: Run integrity_guard.py --fix to rebuild DANH_MUC.html")
    else:
        logger.info("\n[DRY RUN] — No changes made to DB")


if __name__ == "__main__":
    main()
