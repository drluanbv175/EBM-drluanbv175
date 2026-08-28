#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Dựng bản HTML tự chứa TỪ file .docx mà KHÔNG cần pandoc.

VÌ SAO CÓ (12/08/2026): bước ④ của dây chuyền xuất gói cập nhật chứng cứ dựa
hoàn toàn vào `pandoc`. MacBook có pandoc nên ra đủ 5 sản phẩm; máy Windows
KHÔNG có, nên mất luôn cả bước ④ *và* bước ⑤ (PDF giữ màu dựng từ HTML của ④).
Kết quả là hai máy cho ra hai bộ sản phẩm khác nhau từ cùng một dashboard — đúng
thứ mà việc gộp "bộ năm" vào một lệnh sinh ra để tránh.

Bác sĩ chọn hướng bỏ phụ thuộc thay vì bắt cài thêm phần mềm, nên module này
dựng HTML thẳng từ .docx bằng `python-docx` (đã có sẵn trong venv ~/.ebm-venv).

KHÁC BIỆT ĐÁNG CHÚ Ý so với nhánh pandoc: pandoc **bỏ hết màu nền ô**, còn ở đây
ta đọc màu từ chính `w:shd/@w:fill` nên **GIỮ ĐƯỢC** huy hiệu mức chứng cứ và
quyết định. Vì vậy trang sinh bằng nhánh này KHÔNG in cảnh báo "mất màu" — in
cảnh báo đó khi màu vẫn còn là nói sai với người đọc.

Phạm vi có chủ ý: giữ CHỮ · BẢNG · ĐỀ MỤC · THỨ TỰ · MÀU NỀN Ô — đủ để đọc và
để bước ⑤ in ra PDF. KHÔNG tái tạo: ảnh nhúng, hộp văn bản, đánh số tự động của
Word. Bản `.docx` vẫn là bản lưu trữ chuẩn.

Dùng như thư viện:
    from docx_sang_html_khong_pandoc import dung_html_tu_docx
    dung_html_tu_docx(Path("a.docx"), Path("a.html"), tieu_de="...", style=..., banner=...)

Hoặc chạy thẳng:
    python tools/docx_sang_html_khong_pandoc.py <file.docx> [-o out.html]
"""
from __future__ import annotations

import argparse
import html as _html
import pathlib
import sys

# Windows: stdout mặc định là cp1252 → mọi print() tiếng Việt hoặc ký hiệu (✓ ⚠ →)
# ném UnicodeEncodeError và GIẾT tiến trình, thường SAU KHI công việc đã xong. Đo thật
# ngày 12/08/2026 trên dây chuyền cập nhật chứng cứ: bản Word 82 KB đã ghi ra đĩa nhưng
# tool thoát mã 1 ở đúng dòng print cuối ⇒ caller đọc mã thoát, tưởng hỏng, bỏ luôn 2
# bước sau. Cùng lớp lỗi đã vá cho tools/vietnamize/.
import sys as _sys_utf8
for _s in (_sys_utf8.stdout, _sys_utf8.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _thoat(text: str) -> str:
    return _html.escape(text or "", quote=False)


def _mau_nen_o(o) -> str | None:
    """Màu nền của một ô bảng, đọc từ w:shd/@w:fill. None nghĩa là không tô."""
    try:
        tcPr = o._tc.tcPr
        if tcPr is None:
            return None
        shd = tcPr.find(f"{W_NS}shd")
        if shd is None:
            return None
        fill = shd.get(f"{W_NS}fill")
        if not fill or fill in ("auto", "FFFFFF", "ffffff"):
            return None
        return "#" + fill.lstrip("#")
    except Exception:  # noqa: BLE001 — thiếu màu không được làm hỏng cả bản dựng
        return None


def _mau_chu_run(run) -> str | None:
    """Màu CHỮ của một run, đọc từ w:rPr/w:color/@w:val.

    Vì sao cần: bản đầu chỉ đọc màu NỀN ô (w:shd). Với huy hiệu mức chứng cứ và
    quyết định — nền đậm (#15803D xanh · #B45309 cam · #B91C1C đỏ) + chữ TRẮNG —
    bỏ màu chữ làm chữ rơi về đen, tức chữ đen trên nền đỏ đậm. Giữ nền mà mất
    chữ còn khó đọc HƠN là mất cả hai. Đo trên bản Word tâm thần kinh 24/08/2026:
    380 ô có nền, trong đó các ô huy hiệu đều mất chữ trắng.
    """
    try:
        rgb = run.font.color.rgb  # None nếu tự động/theo theme
        if rgb is None:
            return None
        return "#" + str(rgb)
    except Exception:  # noqa: BLE001 — thiếu màu không được làm hỏng cả bản dựng
        return None


def _doan_sang_html(doan) -> str:
    """Một đoạn văn Word → thẻ HTML tương ứng, giữ đề mục và in đậm/nghiêng."""
    ten = (doan.style.name or "").lower() if doan.style is not None else ""
    noi_dung = []
    for run in doan.runs:
        t = _thoat(run.text)
        if not t:
            continue
        if run.bold:
            t = f"<strong>{t}</strong>"
        if run.italic:
            t = f"<em>{t}</em>"
        mau_chu = _mau_chu_run(run)
        if mau_chu:
            t = f'<span style="color:{mau_chu}">{t}</span>'
        noi_dung.append(t)
    text = "".join(noi_dung) or _thoat(doan.text)
    if not text.strip():
        return ""

    if ten.startswith("heading 1") or ten == "title":
        return f"<h1>{text}</h1>"
    if ten.startswith("heading 2"):
        return f"<h2>{text}</h2>"
    if ten.startswith("heading 3"):
        return f"<h3>{text}</h3>"
    if ten.startswith("heading"):
        return f"<h4>{text}</h4>"
    if "list" in ten:
        return f"<li>{text}</li>"
    if "quote" in ten:
        return f"<blockquote>{text}</blockquote>"
    return f"<p>{text}</p>"


def _bang_sang_html(bang) -> str:
    hang_html = []
    for i, hang in enumerate(bang.rows):
        o_html = []
        for o in hang.cells:
            mau = _mau_nen_o(o)
            style = f' style="background:{mau}"' if mau else ""
            noi = " ".join(
                p for p in (_doan_sang_html(d) for d in o.paragraphs) if p
            ) or "&nbsp;"
            the = "th" if i == 0 else "td"
            o_html.append(f"<{the}{style}>{noi}</{the}>")
        hang_html.append("<tr>" + "".join(o_html) + "</tr>")
    return "<table>" + "".join(hang_html) + "</table>"


def _duyet_theo_thu_tu(tai_lieu):
    """Trả về đoạn văn và bảng ĐÚNG THỨ TỰ xuất hiện trong tài liệu.

    python-docx cho `.paragraphs` và `.tables` thành hai danh sách RỜI, ghép lại
    theo kiểu "đoạn trước, bảng sau" sẽ đảo lộn tài liệu. Phải duyệt cây XML gốc.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    than = tai_lieu.element.body
    for con in than.iterchildren():
        if con.tag == f"{W_NS}p":
            yield Paragraph(con, tai_lieu)
        elif con.tag == f"{W_NS}tbl":
            yield Table(con, tai_lieu)


def dung_html_tu_docx(docx_path: pathlib.Path, out_path: pathlib.Path, *,
                      tieu_de: str = "", style: str = "", banner: str = "") -> pathlib.Path:
    """Dựng file HTML tự chứa từ .docx. Trả về đường dẫn đã ghi."""
    import docx
    from docx.table import Table

    tai_lieu = docx.Document(str(docx_path))
    than: list[str] = []
    trong_danh_sach = False

    for khoi in _duyet_theo_thu_tu(tai_lieu):
        if isinstance(khoi, Table):
            if trong_danh_sach:
                than.append("</ul>")
                trong_danh_sach = False
            than.append(_bang_sang_html(khoi))
            continue
        doan_html = _doan_sang_html(khoi)
        if not doan_html:
            continue
        la_muc = doan_html.startswith("<li>")
        if la_muc and not trong_danh_sach:
            than.append("<ul>")
            trong_danh_sach = True
        elif not la_muc and trong_danh_sach:
            than.append("</ul>")
            trong_danh_sach = False
        than.append(doan_html)
    if trong_danh_sach:
        than.append("</ul>")

    tieu_de = tieu_de or docx_path.stem
    trang = (
        "<!doctype html>\n<html lang=\"vi\">\n<head>\n<meta charset=\"utf-8\">\n"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
        f"<title>{_thoat(tieu_de)}</title>\n{style}\n</head>\n<body>\n"
        f"{banner}\n" + "\n".join(than) + "\n</body>\n</html>\n"
    )
    out_path.write_text(trang, encoding="utf-8")
    return out_path


def main() -> int:
    ap = argparse.ArgumentParser(description="Dựng HTML từ .docx không cần pandoc")
    ap.add_argument("docx")
    ap.add_argument("-o", "--out", help="file HTML ra (mặc định cùng tên .html)")
    a = ap.parse_args()
    src = pathlib.Path(a.docx)
    if not src.exists():
        print(f"✗ Không thấy {src}", file=sys.stderr)
        return 2
    out = pathlib.Path(a.out) if a.out else src.with_suffix(".html")
    dung_html_tu_docx(src, out)
    print(f"✓ {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
