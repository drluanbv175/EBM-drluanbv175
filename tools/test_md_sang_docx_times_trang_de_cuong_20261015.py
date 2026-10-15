# -*- coding: utf-8 -*-
"""Chế độ trình bày đề cương của `md_sang_docx_times.py` (15/10/2026).

Bối cảnh: đề cương nộp Hội đồng Y đức (đề tài C1a) cần khổ A4, lề chuẩn, căn đều, ngắt trang trước mỗi chương
và trang bìa căn giữa; bản chuyển đổi thường vẽ «---» thành đường kẻ và để mọi đoạn căn trái. Ngoại tuyến hoàn toàn,
văn bản fixture tự soạn.
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

pytest.importorskip("docx")
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
_sp = importlib.util.spec_from_file_location("_md_docx_dc", REPO / "tools/md_sang_docx_times.py")
M = importlib.util.module_from_spec(_sp)
_sp.loader.exec_module(M)

MD = """**ĐƠN VỊ THỬ**

# ĐỀ CƯƠNG THỬ

## TÊN ĐỀ TÀI THỬ

---

## THÔNG TIN CHUNG

Đoạn văn thân bài thứ nhất.

---

# 1. CHƯƠNG MỘT

- mục danh sách

---

| A | B |
|---|---|
| 1 | 2 |
"""


def _co_ke(p) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.find(qn("w:pBdr")) is not None


def test_mac_dinh_khong_doi():
    doc = M.dung_docx(MD, 12.0)
    assert any(_co_ke(p) for p in doc.paragraphs), "chế độ thường vẫn vẽ đường kẻ"
    assert not any(p.paragraph_format.page_break_before for p in doc.paragraphs)
    assert abs(doc.sections[0].left_margin - Cm(3)) >= Cm(0.01)


def test_trang_de_cuong():
    doc = M.dung_docx(MD, 13.0, trang_de_cuong=True)
    s = doc.sections[0]
    # Word lưu kích thước theo twip (1/1440 inch) ⇒ đọc lại lệch vài trăm EMU so với Cm(...); so trong sai số 0,01 cm.
    gan = lambda x, cm: abs(x - Cm(cm)) < Cm(0.01)  # noqa: E731
    assert gan(s.page_width, 21) and gan(s.page_height, 29.7)
    assert gan(s.left_margin, 3) and gan(s.right_margin, 2) and gan(s.top_margin, 2) and gan(s.bottom_margin, 2)
    assert not any(_co_ke(p) for p in doc.paragraphs), "«---» phải thành ngắt trang, không vẽ kẻ"
    ps = doc.paragraphs
    bia = [p for p in ps if p.text in ("ĐƠN VỊ THỬ", "ĐỀ CƯƠNG THỬ", "TÊN ĐỀ TÀI THỬ")]
    assert len(bia) == 3 and all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in bia)
    ngat = [p.text for p in ps if p.paragraph_format.page_break_before]
    assert ngat[:2] == ["THÔNG TIN CHUNG", "1. CHƯƠNG MỘT"]
    assert len(ngat) == 3 and ngat[2] == "", "bảng ngay sau «---» được đẩy sang trang mới bằng một đoạn ngắt"
    than = next(p for p in ps if p.text.startswith("Đoạn văn thân bài"))
    assert than.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert than.paragraph_format.page_break_before is not True
    chan = s.footer.paragraphs[0]._p.xml
    assert "PAGE" in chan and "fldChar" in chan
